from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


NEW_REFERENCES = [
    "A. F. de Almeida, R. Moreira and T. Rodrigues, Nat. Rev. Chem., 2019, 3, 589-604, DOI: 10.1038/s41570-019-0124-0.",
    "F. Strieth-Kalthoff, F. Sandfort, M. H. S. Segler and F. Glorius, Chem. Soc. Rev., 2020, 49, 6154-6168, DOI: 10.1039/C9CS00786E.",
    "C. W. Coley, N. S. Eyke and K. F. Jensen, Angew. Chem. Int. Ed., 2020, 59, 23414-23436, DOI: 10.1002/anie.201909989.",
    "M. Gutlein, A. Karwath and S. Kramer, J. Cheminform., 2012, 4, 7, DOI: 10.1186/1758-2946-4-7.",
    "S. Szymkuc, E. P. Gajewska, T. Klucznik, K. Molga, P. Dittwald, M. Startek, M. Bajczyk, B. A. Grzybowski, et al., Angew. Chem. Int. Ed., 2016, 55, 5904-5937, DOI: 10.1002/anie.201506101.",
    "V. Venkatasubramanian and V. Mann, Curr. Opin. Chem. Eng., 2022, 36, 100749, DOI: 10.1016/j.coche.2021.100749.",
    "M. Trobe and M. D. Burke, Angew. Chem. Int. Ed., 2018, 57, 4192-4214, DOI: 10.1002/anie.201710482.",
    "M. Christensen, L. P. E. Yunker, P. Shiri, T. Zepel, P. L. Prieto, S. Grunert, F. Bork, J. E. Hein, et al., Chem. Sci., 2021, 12, 15473-15490, DOI: 10.1039/D1SC04588A.",
    "S. Steiner, J. Wolf, S. Glatzel, A. Andreou, J. M. Granda, G. Keenan, T. Hinkley, G. Aragon-Camarasa, P. J. Kitson, D. Angelone and L. Cronin, Science, 2019, 363, eaav2211, DOI: 10.1126/science.aav2211.",
    "N. H. Angello, V. Rathore, W. Beker, A. Wolos, E. R. Jira, R. Roszak, T. C. Wu, C. M. Schroeder, A. Aspuru-Guzik, B. A. Grzybowski and M. D. Burke, Science, 2022, 378, 399-405, DOI: 10.1126/science.adc8743.",
    "R. J. Hickman, M. Aldeghi, F. Hase and A. Aspuru-Guzik, Digit. Discov., 2022, 1, 732-744, DOI: 10.1039/D2DD00028H.",
    "F. Hase, M. Aldeghi, R. J. Hickman, L. M. Roch and A. Aspuru-Guzik, Appl. Phys. Rev., 2021, 8, 031406, DOI: 10.1063/5.0048164.",
    "F. Hase, L. M. Roch and A. Aspuru-Guzik, Chem. Sci., 2018, 9, 7642-7655, DOI: 10.1039/C8SC02239A.",
    "M. Shevlin, ACS Med. Chem. Lett., 2017, 8, 601-607, DOI: 10.1021/acsmedchemlett.7b00165.",
    "V. Sans, L. Porwol, V. Dragone and L. Cronin, Chem. Sci., 2015, 6, 1258-1264, DOI: 10.1039/C4SC03075C.",
    "C. Mateos, M. J. Nieves-Remacha and J. A. Rincon, React. Chem. Eng., 2019, 4, 1536-1544, DOI: 10.1039/C9RE00116F.",
    "J. A. Garrido Torres, S. H. Lau, P. Anchuri, J. M. Stevens, J. E. Tabora, J. Li, A. Borovika, R. P. Adams and A. G. Doyle, J. Am. Chem. Soc., 2022, 144, 19999-20007, DOI: 10.1021/jacs.2c08592.",
    "T. Lookman, P. V. Balachandran, D. Xue and R. Yuan, npj Comput. Mater., 2019, 5, 21, DOI: 10.1038/s41524-019-0153-8.",
    "B. Shahriari, K. Swersky, Z. Wang, R. P. Adams and N. de Freitas, Proc. IEEE, 2016, 104, 148-175, DOI: 10.1109/JPROC.2015.2494218.",
    "D. R. Jones, M. Schonlau and W. J. Welch, J. Glob. Optim., 1998, 13, 455-492, DOI: 10.1023/A:1008306431147.",
]


def find_paragraph(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def replace_citation_run(paragraph, old: str, new: str) -> None:
    matches = [run for run in paragraph.runs if run.text == old and run.font.superscript]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one superscript citation {old!r} in {paragraph.text[:80]!r}, found {len(matches)}"
        )
    matches[0].text = new


def rebuild_paragraph_with_citations(paragraph, segments: list[tuple[str, bool]]) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    for text, superscript in segments:
        run = paragraph.add_run(text)
        if superscript:
            run.font.superscript = True


def append_citation(paragraph, citation: str) -> None:
    text = paragraph.text
    if not text.endswith("."):
        raise RuntimeError(f"Expected paragraph to end with a period: {text[:80]!r}")
    rebuild_paragraph_with_citations(
        paragraph,
        [
            (text[:-1] + " ", False),
            (citation, True),
            (".", False),
        ],
    )


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def integrate(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)

    opening = find_paragraph(document, "Reaction-scope exploration is shifting")
    replace_citation_run(opening, "1-11", "1-11,29-31")

    visualization = find_paragraph(document, "Low-dimensional chemical maps are attractive")
    replace_citation_run(visualization, "12-18", "12-18,32")

    workflows = find_paragraph(document, "This distinction is especially important")
    rebuild_paragraph_with_citations(
        workflows,
        [
            (
                "This distinction is especially important because many automated or "
                "semi-automated chemistry workflows select experiments in descriptor, "
                "fingerprint or model spaces rather than in the final visual display space ",
                False,
            ),
            ("2-11,18-26,33-45", True),
            (
                ". Active learning, Bayesian optimization and diversity sampling can all be "
                "effective, but their behaviour depends on the geometry in which distances, "
                "uncertainty, repulsion or coverage are computed. For reaction-scope studies, "
                "the target is often not only the best yield. Chemists also need to know where "
                "reactivity begins to fail, which functional groups are fragile, and which "
                "apparent holes in a map correspond to genuine unsampled chemistry rather than "
                "projection artifacts.",
                False,
            ),
        ],
    )

    study_design = find_paragraph(document, "The benchmark was designed as a paired")
    replace_citation_run(study_design, "1,12-18,22-28", "1,12-18,22-28,46-48")

    decision_spaces = find_paragraph(document, "Each reaction landscape was represented")
    replace_citation_run(decision_spaces, "12-18", "12-18,32,40")
    replace_citation_run(decision_spaces, "22-24", "22-24")

    multi_objective = find_paragraph(document, "The paper does not assume")
    append_citation(multi_objective, "41,45")

    reporting = find_paragraph(document, "For practical reaction-scope studies")
    append_citation(reporting, "36,41,45")

    candidate_heading = find_paragraph(document, "Candidate Additional References")
    candidate_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is candidate_heading._p
    )
    for paragraph in list(document.paragraphs[candidate_index - 1 :]):
        remove_paragraph(paragraph)

    for reference in NEW_REFERENCES:
        paragraph = document.add_paragraph(reference, style="List Number")
        paragraph.alignment = document.paragraphs[-2].alignment

    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()
    integrate(args.input_docx, args.output_docx)


if __name__ == "__main__":
    main()
