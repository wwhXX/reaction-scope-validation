from __future__ import annotations

import csv
import math
import re
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EXP = ROOT / "experiments"
RESULTS = EXP / "results"
MANUSCRIPT = RESULTS / "manuscript"
REFERENCE = MANUSCRIPT / "reaction_scope_boundary_core_manuscript_chemical_science_revised_20260802.docx"
SOURCE_MD = MANUSCRIPT / "supporting_information_v3.md"
OUTPUT = MANUSCRIPT / "reaction_scope_boundary_supporting_information_v2.docx"

SPRINT = RESULTS / "chemical_science_sprint"
ROBUST = RESULTS / "chemical_science_robustness"
EXTERNAL = RESULTS / "chemical_science_drawable_external"

BLUE = "2E74B5"
NAVY = RGBColor(11, 37, 69)
HEADER_FILL = "F1F3F6"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_font(run, size: float, bold: bool = False, color: RGBColor | None = None, italic: bool = False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=75, bottom=70, end=75):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[float]):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size=8.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(text)
        shade_cell(cell, HEADER_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            set_font(run, font_size, bold=True)
    for values in rows:
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i > 0 and len(str(value)) < 22:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_font(run, font_size)
        for cell in row.cells:
            set_cell_margins(cell)
    for cell in hdr.cells:
        set_cell_margins(cell, top=85, bottom=85)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def metric(value: str, low: str, high: str, digits=4) -> str:
    return f"{float(value):.{digits}f} [{float(low):.{digits}f}, {float(high):.{digits}f}]"


def final_metric_rows(dataset: str, budget: int) -> list[list[str]]:
    rows = read_csv(SPRINT / "decision_space_contrast_r20_summary.csv")
    selected = [r for r in rows if r["dataset"] == dataset and int(float(r["budget_spent"])) == budget]
    order = {"cvt": 0, "weighted_itr_cvt": 1, "diversity": 2, "uncertainty": 3, "uncertainty_diversity": 4}
    dim_order = {"2": 0, "64": 1, "full": 2, "tanimoto": 3}
    selected.sort(key=lambda r: (order.get(r["method"], 99), dim_order.get(r["decision_dimension"], 99)))
    out = []
    for r in selected:
        out.append([
            r["method"].replace("weighted_itr_cvt", "weighted itr. CVT").replace("uncertainty_diversity", "uncertainty + diversity"),
            r["decision_dimension"],
            metric(r["boundary_coverage_mean"], r["boundary_coverage_ci_low"], r["boundary_coverage_ci_high"]),
            metric(r["boundary_enrichment_mean"], r["boundary_enrichment_ci_low"], r["boundary_enrichment_ci_high"], 3),
            metric(r["failure_recall_mean"], r["failure_recall_ci_low"], r["failure_recall_ci_high"]),
            metric(r["failure_cluster_coverage_mean"], r["failure_cluster_coverage_ci_low"], r["failure_cluster_coverage_ci_high"]),
        ])
    return out


def delta_rows() -> list[list[str]]:
    rows = read_csv(SPRINT / "decision_space_contrast_r20_pairwise_deltas.csv")
    rows = [r for r in rows if r["dataset"] in {"aldol", "cobalt"}]
    out = []
    for r in rows:
        out.append([
            r["dataset"],
            r["method"].replace("weighted_itr_cvt", "weighted itr. CVT").replace("uncertainty_diversity", "uncertainty + diversity"),
            r["high_dimension"],
            metric(r["boundary_coverage_delta_mean"], r["boundary_coverage_delta_ci_low"], r["boundary_coverage_delta_ci_high"]),
            metric(r["failure_recall_delta_mean"], r["failure_recall_delta_ci_low"], r["failure_recall_delta_ci_high"]),
            metric(r["failure_cluster_coverage_delta_mean"], r["failure_cluster_coverage_delta_ci_low"], r["failure_cluster_coverage_delta_ci_high"]),
        ])
    return out


def exact_rows() -> list[list[str]]:
    rows = read_csv(MANUSCRIPT / "scopehd_only_boundary_examples.csv")
    out = []
    for r in rows:
        out.append([
            f"{r['row_index']}\n{r['smiles']}",
            f"{float(r['conv']):.0f}\n{'failure' if r['is_failure'].lower() == 'true' else 'success'}",
            f"{r['nearest_success_idx']}\n{r['nearest_success_smiles']}",
            f"{float(r['nearest_success_conv']):.0f}",
            r["highd_only_repeats"],
            r["highd_only_count"],
            f"{float(r['nearest_success_distance']):.1f}",
        ])
    return out


def robustness_grid_rows() -> list[list[str]]:
    rows = read_csv(ROBUST / "boundary_cluster_robustness_r20_pairwise_deltas.csv")
    out = []
    for r in rows:
        out.append([
            r["dataset"], r["high_dimension"], r["boundary_quantile"], r["failure_clusters"],
            metric(r["boundary_coverage_delta_mean"], r["boundary_coverage_delta_ci_low"], r["boundary_coverage_delta_ci_high"]),
            metric(r["failure_recall_delta_mean"], r["failure_recall_delta_ci_low"], r["failure_recall_delta_ci_high"]),
            metric(r["failure_cluster_coverage_delta_mean"], r["failure_cluster_coverage_delta_ci_low"], r["failure_cluster_coverage_delta_ci_high"]),
        ])
    return out


def external_rows() -> list[list[str]]:
    summaries = read_csv(EXTERNAL / "drawable_external_contrast_r20_summary.csv")
    deltas = read_csv(EXTERNAL / "drawable_external_contrast_r20_pairwise_deltas.csv")
    delta_map = {(r["method"], r["high_dimension"]): r for r in deltas}
    out = []
    for r in summaries:
        if r["method"] != "diversity_fps":
            continue
        dim = r["decision_dimension"]
        d = delta_map.get(("diversity_fps", dim))
        out.append([
            dim,
            metric(r["boundary_coverage_mean"], r["boundary_coverage_ci_low"], r["boundary_coverage_ci_high"]),
            metric(r["failure_recall_mean"], r["failure_recall_ci_low"], r["failure_recall_ci_high"]),
            metric(r["failure_cluster_coverage_mean"], r["failure_cluster_coverage_ci_low"], r["failure_cluster_coverage_ci_high"]),
            "baseline" if d is None else metric(d["boundary_coverage_delta_mean"], d["boundary_coverage_delta_ci_low"], d["boundary_coverage_delta_ci_high"]),
            "baseline" if d is None else metric(d["failure_cluster_coverage_delta_mean"], d["failure_cluster_coverage_delta_ci_low"], d["failure_cluster_coverage_delta_ci_high"]),
        ])
    return out


def load_font(size: int, bold=False):
    name = "timesbd.ttf" if bold else "times.ttf"
    path = Path("C:/Windows/Fonts") / name
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


def make_robustness_plot(path: Path):
    data = [
        ("Aldol / full", 0.04332, 0.13018, 0.03874, 0.1390, "#2E74B5"),
        ("Cobalt / full", 0.1214, 0.3167, 0.09286, 0.3722, "#2E8B57"),
        ("Cobalt / Tanimoto", -0.1444, -0.05476, -0.2000, -0.02857, "#B54848"),
    ]
    w, h = 1500, 700
    image = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(image)
    title = load_font(40, True)
    body = load_font(28)
    small = load_font(23)
    draw.text((70, 40), "Boundary-coverage sensitivity ranges", fill="#0B2545", font=title)
    left, right, top, bottom = 380, 1400, 145, 585
    xmin, xmax = -0.22, 0.40
    def x(v):
        return left + (v - xmin) / (xmax - xmin) * (right - left)
    for tick in [-0.2, -0.1, 0, 0.1, 0.2, 0.3, 0.4]:
        xx = x(tick)
        draw.line((xx, top, xx, bottom), fill="#D9DEE5", width=2)
        draw.text((xx - 22, bottom + 12), f"{tick:.1f}", fill="#555555", font=small)
    draw.line((x(0), top - 10, x(0), bottom), fill="#333333", width=4)
    for i, (label, lo, hi, ci_lo, ci_hi, color) in enumerate(data):
        y = 235 + i * 135
        draw.text((70, y - 20), label, fill="#111111", font=body)
        draw.line((x(ci_lo), y, x(ci_hi), y), fill="#333333", width=5)
        draw.line((x(ci_lo), y - 12, x(ci_lo), y + 12), fill="#333333", width=4)
        draw.line((x(ci_hi), y - 12, x(ci_hi), y + 12), fill="#333333", width=4)
        draw.line((x(lo), y, x(hi), y), fill=color, width=18)
        draw.ellipse((x(lo)-9, y-9, x(lo)+9, y+9), fill=color)
        draw.ellipse((x(hi)-9, y-9, x(hi)+9, y+9), fill=color)
    draw.text((left + 230, 650), "Paired richer-space-minus-2D boundary-coverage difference", fill="#333333", font=small)
    image.save(path)


FIGURES = {
    "Figure S1.": (MANUSCRIPT / "scopehd_vs_2d_boundary_comparison.png", 7.0),
    "Figure S2.": (SPRINT / "failure_cluster_explanations_aldol_structures.png", 6.45),
    "Figure S3.": (SPRINT / "failure_cluster_explanations_cobalt_structures.png", 7.0),
    "Figure S4.": (MANUSCRIPT / "boundary_robustness_ranges_si.png", 7.0),
    "Figure S5.": (EXTERNAL / "drawable_external_contrast_r20_representatives_structures.png", 5.65),
}


DYNAMIC_TABLES = {
    "Table S3.": (["Method", "Space", "Boundary coverage", "Boundary enrichment", "Failure recall", "Failure-cluster coverage"], lambda: final_metric_rows("aldol", 80), [1.08, 0.55, 1.47, 1.32, 1.30, 1.48], 7.2),
    "Table S4.": (["Method", "Space", "Boundary coverage", "Boundary enrichment", "Failure recall", "Failure-cluster coverage"], lambda: final_metric_rows("cobalt", 24), [1.08, 0.55, 1.47, 1.32, 1.30, 1.48], 7.2),
    "Table S5.": (["Dataset", "Method", "Richer space", "Boundary coverage delta", "Failure recall delta", "Failure-cluster delta"], delta_rows, [0.65, 1.15, 0.70, 1.55, 1.55, 1.60], 7.0),
    "Table S6.": (["Miss row / SMILES", "Target / status", "Nearest-success row / SMILES", "Success target", "Full-only repeats", "n", "Distance"], exact_rows, [1.43, 0.57, 1.43, 0.48, 1.28, 0.32, 0.49], 6.5),
    "Table S9.": (["Dataset", "Space", "q", "K", "Boundary delta", "Failure-recall delta", "Failure-cluster delta"], robustness_grid_rows, [0.62, 0.65, 0.35, 0.35, 1.65, 1.65, 1.65], 6.8),
    "Table S10.": (["Space", "Boundary coverage", "Failure recall", "Failure-cluster coverage", "Boundary delta", "Cluster delta"], external_rows, [0.55, 1.47, 1.36, 1.47, 1.47, 1.40], 7.1),
}


def clean_inline(text: str) -> str:
    return text.replace("`", "").replace("**", "")


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    match = re.match(r"((?:Figure|Table) S\d+\.)(.*)", text)
    if match:
        r1 = p.add_run(match.group(1))
        set_font(r1, 9, bold=True)
        r2 = p.add_run(match.group(2))
        set_font(r2, 9)
    else:
        run = p.add_run(text)
        set_font(run, 9)


def parse_markdown_table(lines: list[str], index: int):
    rows = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [clean_inline(c.strip()) for c in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        index += 1
    return rows[0], rows[1:], index


def build():
    make_robustness_plot(MANUSCRIPT / "boundary_robustness_ranges_si.png")
    work = MANUSCRIPT / ".si_reference_working_copy.docx"
    shutil.copy2(REFERENCE, work)
    doc = Document(work)
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), "1")

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.125
    for style_name, size, before, after in (("Heading 1", 13, 12, 6), ("Heading 2", 10.5, 9, 3)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(46, 116, 181)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.line_spacing = 1.0

    kicker = doc.add_paragraph()
    r = kicker.add_run("Supporting Information")
    set_font(r, 13, bold=True, color=RGBColor(46, 116, 181))
    kicker.paragraph_format.space_after = Pt(6)
    title = doc.add_paragraph()
    r = title.add_run("Visualization Space Is Not Validation Space: A Focused Boundary Benchmark for Reaction-Scope Mapping")
    set_font(r, 20, bold=True, color=NAVY)
    title.paragraph_format.space_after = Pt(10)
    author = doc.add_paragraph()
    set_font(author.add_run("Wenhao Wang"), 11)
    sup = author.add_run("1")
    set_font(sup, 8)
    sup.font.superscript = True
    set_font(author.add_run(" and Jun Zhu"), 11)
    sup = author.add_run("*2")
    set_font(sup, 8)
    sup.font.superscript = True
    author.paragraph_format.space_after = Pt(6)
    aff1 = doc.add_paragraph()
    sup = aff1.add_run("1")
    set_font(sup, 8)
    sup.font.superscript = True
    set_font(aff1.add_run(" Guangdong Basic Research Center of Excellence for Aggregate Science, School of Science and Engineering, The Chinese University of Hong Kong (Shenzhen), Longgang, Shenzhen, Guangdong, 518172, P.R. China."), 10)
    aff1.paragraph_format.space_after = Pt(3)
    aff2 = doc.add_paragraph()
    sup = aff2.add_run("2")
    set_font(sup, 8)
    sup.font.superscript = True
    set_font(aff2.add_run(" School of Science and Engineering, The Chinese University of Hong Kong, Shenzhen, Guangdong 518172, China"), 10)
    aff2.paragraph_format.space_after = Pt(3)
    corr = doc.add_paragraph()
    set_font(corr.add_run("Corresponding-author e-mail: [to be supplied by the authors]"), 9.5, italic=True)
    corr.paragraph_format.space_after = Pt(10)

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## Scope of this Supporting Information"))
    i = start
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            text = clean_inline(line[3:])
            if text.startswith("S1."):
                doc.add_page_break()
            doc.add_heading(text, level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=2)
            i += 1
            continue
        if line == "\\[":
            eq = []
            i += 1
            while i < len(lines) and lines[i].strip() != "\\]":
                eq.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            set_font(p.add_run(" ".join(eq)), 10, italic=True)
            i += 1
            continue
        figure_key = next((key for key in FIGURES if line.startswith(key)), None)
        if figure_key:
            path, width = FIGURES[figure_key]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Inches(width))
            add_caption(doc, clean_inline(line))
            i += 1
            continue
        table_key = re.match(r"Table S\d+\.", line)
        if table_key:
            key = table_key.group(0)
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and lines[j].strip().startswith("|"):
                headers, rows, next_i = parse_markdown_table(lines, j)
                widths = [7.2 / len(headers)] * len(headers)
                add_table(doc, headers, rows, widths, font_size=7.8 if len(headers) > 5 else 8.2)
                add_caption(doc, clean_inline(line))
                i = next_i
                continue
            if key in DYNAMIC_TABLES:
                headers, row_fn, widths, font_size = DYNAMIC_TABLES[key]
                add_table(doc, headers, row_fn(), widths, font_size=font_size)
                add_caption(doc, clean_inline(line))
                i += 1
                continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            set_font(p.add_run(clean_inline(line[2:])), 9.5)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue
        p = doc.add_paragraph()
        set_font(p.add_run(clean_inline(line)), 9.5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1

    footer = section.footer.paragraphs[0]
    footer.text = ""

    doc.core_properties.title = "Supporting Information - Visualization Space Is Not Validation Space"
    doc.core_properties.author = "Wenhao Wang; Jun Zhu"
    doc.core_properties.subject = "Supporting Information for the reaction-scope decision-space benchmark"
    doc.save(OUTPUT)
    work.unlink(missing_ok=True)
    print(OUTPUT)


if __name__ == "__main__":
    build()
