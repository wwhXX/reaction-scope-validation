# -*- coding: utf-8 -*-
"""Build a pruned core DOCX manuscript from the Chemical Science sprint outputs."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SPRINT_DIR = ROOT / "experiments" / "results" / "chemical_science_sprint"
ROBUSTNESS_DIR = ROOT / "experiments" / "results" / "chemical_science_robustness"
DRAWABLE_DIR = ROOT / "experiments" / "results" / "chemical_science_drawable_external"
OUT_DIR = ROOT / "experiments" / "results" / "manuscript"
DOCX_OUT = OUT_DIR / "reaction_scope_boundary_core_manuscript.docx"
DOCX_FALLBACK_OUT = OUT_DIR / "reaction_scope_boundary_core_manuscript_chemical_science.docx"
CORE_DELTA_FIGURE = OUT_DIR / "core_aldol_cobalt_boundary_deltas.png"
PAIRED_MISS_SOURCE_FIGURE = OUT_DIR / "scopehd_vs_2d_boundary_comparison.png"
PAIRED_MISS_MAIN_FIGURE = OUT_DIR / "scopehd_vs_2d_boundary_comparison_main.png"

def set_section_columns(section, count: int = 1, space_dxa: int = 360) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_dxa))


def switch_columns(doc: Document, count: int) -> None:
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_section_columns(section, count=count)


def add_full_width_block(doc: Document, callback) -> None:
    switch_columns(doc, 1)
    callback()
    switch_columns(doc, 2)


def justify_last_paragraph(doc: Document) -> None:
    if doc.paragraphs:
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{margin_name}"))
        if element is None:
            element = OxmlElement(f"w:{margin_name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)


def format_table(table, widths: list[int]) -> None:
    set_table_geometry(table, widths)
    for idx, row in enumerate(table.rows):
        if idx == 0:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = tr_pr.find(qn("w:tblHeader"))
            if tbl_header is None:
                tbl_header = OxmlElement("w:tblHeader")
                tr_pr.append(tbl_header)
            tbl_header.set(qn("w:val"), "true")
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(0)
            if idx == 0:
                set_cell_shading(cell, "F4F6F9")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    set_section_columns(section, count=1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.125
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in [
        ("Heading 1", 13, RGBColor(46, 116, 181), 12, 6),
        ("Heading 2", 10.5, RGBColor(46, 116, 181), 9, 3),
        ("Heading 3", 10, RGBColor(31, 77, 120), 7, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    match = re.match(r"^((?:Figure|Table)\s+\d+\.)\s*(.*)$", text)
    if match:
        label = paragraph.add_run(match.group(1) + " ")
        label.bold = True
        body = paragraph.add_run(match.group(2))
        body.bold = False
    else:
        run = paragraph.add_run(text)
        run.bold = False


def add_paragraph_with_citations(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    if style is None:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cursor = 0
    for match in re.finditer(r"\[([0-9,\-]+)\]", text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        run = paragraph.add_run(match.group(1))
        run.font.superscript = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])
    return paragraph


def add_note(doc: Document, label: str, text: str) -> None:
    switch_columns(doc, 1)
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_text(cell, f"{label}: {text}")
    set_cell_shading(cell, "EAF2F8")
    set_table_geometry(table, [9360])
    switch_columns(doc, 2)


def add_picture_if_exists(doc: Document, path: Path, caption: str, width: float = 6.25) -> None:
    if not path.exists():
        return
    switch_columns(doc, 1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)
    switch_columns(doc, 2)


def build_core_delta_figure() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    deltas = pd.read_csv(SPRINT_DIR / "chemical_science_key_boundary_deltas.csv")
    deltas = deltas[deltas["dataset"].isin(["aldol", "cobalt"])].copy()
    deltas["label"] = deltas.apply(
        lambda row: f"{row['dataset']} | {row['method']} | {row['high_dimension_vs_2d']}",
        axis=1,
    )
    deltas = deltas.sort_values("boundary_coverage_delta_mean", ascending=True)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    font = ImageFont.truetype(str(font_path), 24) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 20) if font_path.exists() else ImageFont.load_default()
    title_font = ImageFont.truetype(str(bold_path), 34) if bold_path.exists() else font

    width, height = 1600, 820
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 45), "Frozen main evidence: Aldol and Cobalt paired r20 contrasts", fill="#0B2545", font=title_font)
    draw.text((80, 96), "Final-budget boundary-coverage delta over 2D, with bootstrap 95% CI", fill="#555555", font=small)

    plot_left, plot_right = 630, 1490
    plot_top, row_h = 175, 90
    axis_min, axis_max = 0.0, 0.55
    axis_y = plot_top + row_h * len(deltas) + 30
    draw.line((plot_left, axis_y, plot_right, axis_y), fill="#333333", width=2)
    for tick in [0.0, 0.1, 0.2, 0.3, 0.4, 0.5]:
        x = plot_left + int((tick - axis_min) / (axis_max - axis_min) * (plot_right - plot_left))
        draw.line((x, plot_top - 18, x, axis_y), fill="#E6EAF0", width=1)
        draw.text((x - 18, axis_y + 12), f"{tick:.1f}", fill="#555555", font=small)

    for idx, row in enumerate(deltas.itertuples(index=False)):
        y = plot_top + idx * row_h
        value = float(row.boundary_coverage_delta_mean)
        ci_low = float(row.boundary_coverage_delta_ci_low)
        ci_high = float(row.boundary_coverage_delta_ci_high)
        x_value = plot_left + int((value - axis_min) / (axis_max - axis_min) * (plot_right - plot_left))
        x_low = plot_left + int((ci_low - axis_min) / (axis_max - axis_min) * (plot_right - plot_left))
        x_high = plot_left + int((ci_high - axis_min) / (axis_max - axis_min) * (plot_right - plot_left))
        color = "#2E8B57" if row.dataset == "cobalt" else "#4F81BD"
        label = f"{row.dataset} | {row.method} | {row.high_dimension_vs_2d}"
        draw.text((80, y + 18), label, fill="#111111", font=font)
        draw.rounded_rectangle((plot_left, y + 18, x_value, y + 52), radius=6, fill=color)
        draw.line((x_low, y + 35, x_high, y + 35), fill="#222222", width=3)
        draw.line((x_low, y + 25, x_low, y + 45), fill="#222222", width=3)
        draw.line((x_high, y + 25, x_high, y + 45), fill="#222222", width=3)
        draw.text((x_value + 14, y + 17), f"{value:.4f}", fill="#111111", font=small)

    draw.text((plot_left, height - 80), "Boundary-coverage delta over 2D", fill="#333333", font=small)
    image.save(CORE_DELTA_FIGURE)


def build_main_paired_miss_figure() -> None:
    """Prune the full exact-miss panel into a main-text-sized Figure 2."""
    if not PAIRED_MISS_SOURCE_FIGURE.exists():
        return

    source = Image.open(PAIRED_MISS_SOURCE_FIGURE).convert("RGB")
    canvas = Image.new("RGB", (1800, 1180), "white")
    draw = ImageDraw.Draw(canvas)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    title_font = ImageFont.truetype(str(bold_path), 36) if bold_path.exists() else ImageFont.load_default()
    label_font = ImageFont.truetype(str(bold_path), 24) if bold_path.exists() else ImageFont.load_default()
    body_font = ImageFont.truetype(str(font_path), 20) if font_path.exists() else ImageFont.load_default()
    small_font = ImageFont.truetype(str(font_path), 17) if font_path.exists() else ImageFont.load_default()

    draw.text((50, 32), "Full-space vs 2D: exact paired-miss boundary examples", fill="#0B2545", font=title_font)
    draw.text(
        (50, 80),
        "Structure pairs are enlarged for main-text reading; the 2D projection is retained as a location index.",
        fill="#555555",
        font=body_font,
    )

    structure_panel = source.crop((45, 68, 1150, 850)).resize((1240, 890), Image.Resampling.LANCZOS)
    canvas.paste(structure_panel, (40, 135))

    inset_x, inset_y = 1285, 165
    draw.text((inset_x, inset_y), "B. 2D location index", fill="#0B2545", font=label_font)
    projection = source.crop((1215, 170, 2320, 805)).resize((470, 270), Image.Resampling.LANCZOS)
    canvas.paste(projection, (inset_x, inset_y + 42))
    draw.rounded_rectangle((inset_x, inset_y + 42, inset_x + 470, inset_y + 312), radius=12, outline="#D7DEE8", width=2)
    draw.text((inset_x, inset_y + 330), "Red rings: HD-selected / 2D-missed rows", fill="#555555", font=small_font)
    draw.text((inset_x, inset_y + 356), "Blue rings: nearest high-D successes", fill="#555555", font=small_font)

    summary_y = 610
    draw.text((inset_x, summary_y), "Row-level replay check", fill="#0B2545", font=label_font)
    draw.text((inset_x, summary_y + 48), "Mean boundary rows per repeat:", fill="#333333", font=body_font)
    draw.rounded_rectangle((inset_x, summary_y + 88, inset_x + 245, summary_y + 118), radius=6, fill="#8FBCE6")
    draw.text((inset_x + 260, summary_y + 84), "2D: 19.0", fill="#333333", font=body_font)
    draw.rounded_rectangle((inset_x, summary_y + 142, inset_x + 360, summary_y + 172), radius=6, fill="#2E8B57")
    draw.text((inset_x + 15, summary_y + 179), "full-space: 37.9", fill="#333333", font=body_font)
    draw.text((inset_x, summary_y + 215), "Mean full-only boundary rows: 30.0", fill="#B54848", font=body_font)
    draw.text((inset_x, summary_y + 255), "The r20 paired benchmark remains", fill="#555555", font=small_font)
    draw.text((inset_x, summary_y + 280), "the statistical result.", fill="#555555", font=small_font)

    draw.text(
        (50, 1128),
        (
            "Aldol row-level paired replay; examples are selected in full descriptor space and missed by the paired 2D replay."
        ),
        fill="#666666",
        font=small_font,
    )
    canvas.save(PAIRED_MISS_MAIN_FIGURE)


def add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(
        "Visualization Space Is Not Validation Space: "
        "A Focused Boundary Benchmark for Reaction-Scope Mapping"
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Core manuscript draft after experiment pruning")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(89, 89, 89)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    run = meta.add_run("Frozen evidence package: Aldol and Cobalt main text; robustness and drawable external controls as support")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_abstract(doc: Document) -> None:
    doc.add_heading("Abstract", level=1)
    add_paragraph_with_citations(
        doc,
        "Reaction-scope maps are increasingly used to guide how chemists choose the next substrates to test, "
        "especially in human-in-the-loop and data-driven workflows [1-11]. Yet the two-dimensional picture that "
        "helps a chemist discuss a scope is not necessarily a validated space for acquisition. Here we treat that "
        "distinction as an empirical question. Using the ScopeMap aldol and cobalt reaction-scope landscapes, we "
        "ran paired retrospective campaigns in which the initial sets, budgets, and scoring geometry were fixed, "
        "while the acquisition space was varied between 2D projections and chemically richer descriptor or "
        "fingerprint spaces [1,12-18,22-24]. In both systems, low-dimensional acquisition missed reproducible "
        "near-boundary regions. At the final budget, the strongest cobalt contrast increased boundary coverage by "
        "0.4867 over 2D, and the strongest aldol contrast by 0.1293; the corresponding bootstrap intervals remained "
        "positive [27]. The missed points were not random scatter: they included nitro/electron-withdrawing, "
        "alkoxy/benzyloxy, sulfonyl/pyridyl, and donor-substituted aldehyde regions in aldol, and "
        "condition-dependent zero-yield alcohol regions in cobalt. Sensitivity analyses over boundary definitions "
        "and failure-cluster counts supported the same direction, while a drawable Buchwald-Hartwig control is kept "
        "as external support rather than as the headline result. These results suggest a simple reporting rule: "
        "2D maps can remain useful interfaces, but acquisition and boundary-discovery claims should be validated "
        "in the decision spaces that generated them."
    )


def add_introduction(doc: Document) -> None:
    doc.add_heading("Introduction", level=1)
    intro_paragraphs = [
        (
            "Reaction-scope exploration is shifting from a static reporting exercise to a decision-support problem. "
            "In contemporary data-driven and human-in-the-loop workflows, maps of substrate space are used not only "
            "to summarize what has already been tested, but also to help chemists decide which substrates, components "
            "or conditions should be tested next [1-11]. This change gives reaction-scope maps a dual role: they must "
            "remain chemically legible to human readers, while also supporting acquisition decisions that determine "
            "which regions of the experimental landscape are observed."
        ),
        (
            "Low-dimensional chemical maps are attractive because they make large substrate spaces inspectable. "
            "Principal-component, manifold-learning and fingerprint-based visualizations can expose clusters, trends "
            "and apparent gaps in chemical space [12-18]. However, visual clarity is not the same as acquisition "
            "validity. A projection that is useful for discussion may compress steric, electronic or functional-group "
            "variation that is essential for locating the boundary between compatible, marginal and failed substrates. "
            "The practical question is therefore not whether two-dimensional maps are useful, but whether they should "
            "be treated as validated decision spaces when the objective is boundary discovery."
        ),
        (
            "This distinction is especially important because many automated or semi-automated chemistry workflows "
            "select experiments in descriptor, fingerprint or model spaces rather than in the final visual display "
            "space [2-11,18-26]. Active learning, Bayesian optimization and diversity sampling can all be effective, but "
            "their behaviour depends on the geometry in which distances, uncertainty, repulsion or coverage are "
            "computed. For reaction-scope studies, the target is often not only the best yield. Chemists also need to "
            "know where reactivity begins to fail, which functional groups are fragile, and which apparent holes in "
            "a map correspond to genuine unsampled chemistry rather than projection artifacts."
        ),
        (
            "Recent reaction-space and substrate-scope studies have made major progress in representing reactions, "
            "predicting reaction performance and designing iterative experimental campaigns [1-11,18-21]. Yet there is "
            "still a methodological gap: acquisition spaces are rarely audited with metrics designed for reaction "
            "boundaries. Standard predictive accuracy or visual coverage can miss the question that matters for a "
            "scope map: under the same experimental budget, does one decision space recover more chemically "
            "meaningful near-boundary and failure-prone regions than another?"
        ),
        (
            "Here we treat this as an empirical benchmark problem. Using the published ScopeMap aldol and cobalt "
            "reaction-scope landscapes as chemically interpretable case studies, we perform paired retrospective "
            "campaigns in which the initial set, sampling budget and final scoring geometry are fixed, while the "
            "acquisition space is varied between two-dimensional projections and chemically richer descriptor or "
            "fingerprint spaces [1,12-18,22-24]. We evaluate final-budget boundary coverage, failure recall, failure-cluster "
            "coverage and boundary enrichment, and we connect the metric outcomes to representative substrate "
            "families. The resulting claim is deliberately narrow: two-dimensional maps can remain valuable visual "
            "interfaces, but boundary-discovery claims should be validated in the decision spaces that generated the "
            "experiments."
        ),
    ]
    for text in intro_paragraphs:
        add_paragraph_with_citations(doc, text)


def add_references(doc: Document) -> None:
    doc.add_heading("References", level=1)
    references = [
        "J. Li, X. Xiao, Q. Yang, B. Zhao and S. Luo, Angew. Chem. Int. Ed., 2026, 65, e2455429, DOI: 10.1002/anie.2455429.",
        "D. T. Ahneman, J. G. J. Estrada, S. Lin, S. D. Dreher and A. G. Doyle, Science, 2018, 360, 186-190, DOI: 10.1126/science.aar5169.",
        "C. W. Coley, N. S. Eyke and K. F. Jensen, Angew. Chem. Int. Ed., 2020, 59, 22858-22893, DOI: 10.1002/anie.201909987.",
        "B. J. Shields, J. Stevens, J. Li, M. Parasram, F. Damani, J. I. Martinez Alvarado, J. M. Janey, R. P. Adams and A. G. Doyle, Nature, 2021, 590, 89-96, DOI: 10.1038/s41586-021-03213-y.",
        "J. M. Granda, L. Donina, V. Dragone, D.-L. Long and L. Cronin, Nature, 2018, 559, 377-381, DOI: 10.1038/s41586-018-0307-8.",
        "B. Burger, P. M. Maffettone, V. V. Gusev, C. M. Aitchison, Y. Bai, X. Wang, X. Li, B. M. Alston, B. Li, R. Clowes, N. Rankin, B. Harris, R. S. Sprick and A. I. Cooper, Nature, 2020, 583, 237-241, DOI: 10.1038/s41586-020-2442-2.",
        "C. W. Coley, D. A. Thomas III, J. A. M. Lummiss, J. N. Jaworski, C. P. Breen, V. Schultz, T. Hart, J. S. Fishman, L. Rogers, H. Gao, R. W. Hicklin, P. P. Plehiers, J. Byington, J. S. Piotti, W. H. Green, A. J. Hart, T. F. Jamison and K. F. Jensen, Science, 2019, 365, eaax1566, DOI: 10.1126/science.aax1566.",
        "F. Hase, L. M. Roch and A. Aspuru-Guzik, Trends Chem., 2019, 1, 282-291, DOI: 10.1016/j.trechm.2019.02.007.",
        "F. Hase, L. M. Roch, C. Kreisbeck and A. Aspuru-Guzik, ACS Cent. Sci., 2018, 4, 1134-1145, DOI: 10.1021/acscentsci.8b00307.",
        "A. F. Zahrt, J. J. Henle, B. T. Rose, Y. Wang, W. T. Darrow and S. E. Denmark, Science, 2019, 363, eaau5631, DOI: 10.1126/science.aau5631.",
        "F. Sandfort, F. Strieth-Kalthoff, M. Kuhnemund, C. Beecks and F. Glorius, Chem, 2020, 6, 1379-1390, DOI: 10.1016/j.chempr.2020.02.017.",
        "H. L. Morgan, J. Chem. Doc., 1965, 5, 107-113, DOI: 10.1021/c160017a018.",
        "D. Rogers and M. Hahn, J. Chem. Inf. Model., 2010, 50, 742-754, DOI: 10.1021/ci100050t.",
        "D. Bajusz, A. Racz and K. Heberger, J. Cheminform., 2015, 7, 20, DOI: 10.1186/s13321-015-0069-3.",
        "L. McInnes, J. Healy and J. Melville, J. Open Source Softw., 2018, 3, 861, DOI: 10.21105/joss.00861.",
        "L. van der Maaten and G. Hinton, J. Mach. Learn. Res., 2008, 9, 2579-2605.",
        "D. Probst and J.-L. Reymond, J. Cheminform., 2020, 12, 12, DOI: 10.1186/s13321-020-0416-x.",
        "P. Schwaller, D. Probst, A. C. Vaucher, V. H. Nair, D. Kreutter, T. Laino and J.-L. Reymond, Nat. Mach. Intell., 2021, 3, 144-152, DOI: 10.1038/s42256-020-00284-w.",
        "P. Schwaller, T. Laino, T. Gaudin, P. Bolgar, C. A. Hunter, C. Bekas and A. A. Lee, ACS Cent. Sci., 2019, 5, 1572-1583, DOI: 10.1021/acscentsci.9b00576.",
        "M. H. S. Segler, M. Preuss and M. P. Waller, Nature, 2018, 555, 604-610, DOI: 10.1038/nature25978.",
        "J. N. Wei, D. Duvenaud and A. Aspuru-Guzik, ACS Cent. Sci., 2016, 2, 725-732, DOI: 10.1021/acscentsci.6b00219.",
        "Q. Du, V. Faber and M. Gunzburger, SIAM Rev., 1999, 41, 637-676, DOI: 10.1137/S0036144599352836.",
        "R. W. Kennard and L. A. Stone, Technometrics, 1969, 11, 137-148, DOI: 10.1080/00401706.1969.10490666.",
        "M. D. McKay, R. J. Beckman and W. J. Conover, Technometrics, 1979, 21, 239-245, DOI: 10.1080/00401706.1979.10489755.",
        "B. Settles, Active Learning Literature Survey, University of Wisconsin-Madison, Madison, WI, 2009.",
        "J. Snoek, H. Larochelle and R. P. Adams, Adv. Neural Inf. Process. Syst., 2012, 25, 2951-2959.",
        "B. Efron and R. J. Tibshirani, An Introduction to the Bootstrap, Chapman and Hall/CRC, New York, 1993.",
        "D. C. Montgomery, Design and Analysis of Experiments, Wiley, Hoboken, 10th edn, 2019.",
    ]
    for item in references:
        paragraph = doc.add_paragraph(item, style="List Number")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_key_message(doc: Document) -> None:
    doc.add_heading("Core Decision", level=1)
    add_note(
        doc,
        "Manuscript scope",
        "The main text is frozen around Aldol and Cobalt paired decision-space contrasts. Robustness, drawable "
        "external Buchwald-Hartwig support, reference-region recovery, Suzuki-Miyaura, and broad metric-tradeoff "
        "analyses belong in the supplement or limitations, not as competing headline stories.",
    )
    doc.add_paragraph(
        "This pruning changes the paper from a catalogue of many experiments into a short, testable claim: a "
        "two-dimensional map can be a good visual summary while still being a poor acquisition space for boundary "
        "exploration."
    )


def add_methods(doc: Document) -> None:
    doc.add_heading("Experimental", level=1)
    doc.add_heading("Study design", level=2)
    add_paragraph_with_citations(
        doc,
        "The benchmark was designed as a paired retrospective acquisition study rather than as a new "
        "reaction-prediction model. The question was whether the geometry used to choose new scope entries "
        "changes which reaction-boundary regions are recovered under the same experimental budget. For each "
        "reaction landscape, we fixed the labelled dataset, target definition, initial set, acquisition budget "
        "and final scoring space, then varied only the decision space used by the sampling algorithm. This design "
        "keeps the comparison focused on the manuscript claim: a two-dimensional map can be a useful visual "
        "interface without being a validated acquisition space [1,12-18,22-28]."
    )

    doc.add_heading("Reaction-scope datasets", level=2)
    doc.add_paragraph(
        "Two ScopeMap case studies were retained as the main text evidence. The aldol dataset contains 1091 "
        "labelled aldehyde entries with conversion as the response variable. Entries with conversion at or above "
        "70% were treated as successful for binary failure and boundary analyses. The cobalt dataset contains 60 "
        "labelled alcohol entries from condition 1 of the cobalt-catalysed transformation. Because this dataset "
        "is small and contains many zero-yield entries, any nonzero condition-1 yield was treated as successful, "
        "and the system is used as small-data support rather than as the sole budget-generalization claim."
    )

    doc.add_heading("Frozen evidence hierarchy", level=2)
    doc.add_paragraph(
        "The frozen evidence hierarchy keeps aldol and cobalt in the main text because both retain chemically "
        "interpretable substrate structures and support paired two-dimensional versus higher-dimensional "
        "acquisition contrasts. Boundary-quantile sensitivity, failure-cluster-count sensitivity and the drawable "
        "Buchwald-Hartwig benchmark are retained as supporting evidence. Suzuki-Miyaura, article reference-region "
        "recovery and broad sampler-ranking scans are treated as supplementary or future-work material because "
        "they broaden the paper without strengthening the central claim."
    )
    switch_columns(doc, 1)
    table = doc.add_table(rows=1, cols=4)
    headers = ["Layer", "Content", "Role", "Manuscript placement"]
    for idx, header in enumerate(headers):
        set_cell_text(table.cell(0, idx), header, bold=True)
    rows = [
        ("Main", "Aldol and Cobalt r20 paired decision-space contrast", "Tests the central 2D-vs-decision-space claim", "Main Results"),
        ("Main", "Aldol/Cobalt chemically reviewed failure-region labels", "Converts metric gains into chemistry-facing regions", "Main Results"),
        ("Support", "Boundary-quantile and failure-cluster robustness", "Stress-tests boundary definition choices", "Supplement / short main note"),
        ("Support", "Drawable Buchwald-Hartwig control", "Shows external full-label structure handling", "Supplement / limitations"),
        ("Archive", "Reference-region, Suzuki, broad Pareto/tradeoff scans", "Useful background, but not needed for the core claim", "Supplement or future work"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    format_table(table, [1100, 3100, 2700, 2460])
    switch_columns(doc, 2)

    doc.add_heading("Decision spaces and acquisition policies", level=2)
    add_paragraph_with_citations(
        doc,
        "Each reaction landscape was represented in a two-dimensional visualization space and in chemically richer "
        "descriptor or fingerprint spaces. The two-dimensional space represents the kind of projected map that is "
        "useful for human inspection. Higher-dimensional descriptor spaces preserve more of the original "
        "molecular-feature geometry, while Tanimoto space uses fingerprint similarity to compare candidate "
        "substrates [12-18]. The main paired contrasts used CVT-style and weighted iterative CVT-style acquisition, "
        "with Kennard-Stone and related diversity baselines retained for context [22-24]. These procedures are "
        "interpreted as geometric acquisition policies, not as mechanistic models of reactivity."
    )
    add_paragraph_with_citations(
        doc,
        "For each repeat, the same initial selected set was shared across decision spaces. Candidate selection was "
        "then performed in the assigned decision space, but final evaluation was always performed in a fixed "
        "full-space scoring geometry. This paired design isolates the effect of decision-space choice from random "
        "initial-set differences. Aldol campaigns used an initial set of 10 selected entries and a final selected "
        "budget of 80. Cobalt campaigns used an initial set of 6 selected entries and a final selected budget of "
        "24. Cobalt therefore tests whether the same failure-region logic is visible in a constrained small-data "
        "setting, while aldol carries the broader boundary-discovery argument."
    )

    doc.add_heading("Boundary-aware endpoints", level=2)
    doc.add_paragraph(
        "The primary endpoint is final-budget boundary coverage: the fraction of full-space near-boundary points "
        "sampled by the campaign. Secondary endpoints include boundary enrichment, failure recall, and "
        "failure-cluster coverage. These metrics intentionally separate several chemistry-facing questions: "
        "whether the campaign reaches the decision boundary, whether it preferentially enriches boundary points, "
        "whether it finds failed substrates, and whether those failures span distinct failure regions."
    )
    doc.add_paragraph(
        "The paper does not assume that a single decision space should dominate every endpoint. A method can "
        "increase boundary coverage while decreasing failure-cluster coverage or failure recall. Such conflicts "
        "are treated as part of the result, because reaction-scope mapping is a multi-objective acquisition "
        "problem rather than a single-score ranking problem."
    )

    doc.add_heading("Statistical summaries and robustness checks", level=2)
    add_paragraph_with_citations(
        doc,
        "The main results use 20 paired repeats for the aldol and cobalt decision-space contrasts. Reported "
        "deltas compare each higher-dimensional acquisition space against the matched two-dimensional acquisition "
        "baseline at the same final budget. Confidence intervals are bootstrap summaries over paired repeats [27]. The "
        "main text emphasizes contrasts whose intervals remain positive for final-budget boundary coverage and "
        "then checks whether the newly sampled regions are chemically structured rather than random scatter."
    )
    doc.add_paragraph(
        "Robustness checks varied the boundary quantile and the number of failure clusters while preserving "
        "shared initial sets and fixed full-space evaluation. These checks are reported as sensitivity analyses "
        "because the robustness implementation used a NumPy-only weighted-iterative-CVT-like approximation to "
        "avoid local numerical-library instability. The drawable Buchwald-Hartwig control was used to verify that "
        "the workflow can retain component-level structures in an external labelled reaction dataset; it is not "
        "used as the headline result because its boundary-coverage behaviour does not provide the strongest "
        "high-dimensional advantage."
    )

    doc.add_heading("Chemistry-facing interpretation", level=2)
    doc.add_paragraph(
        "Failure-region labels were assigned after the numerical benchmark to make the missed regions chemically "
        "readable. In aldol, the retained low-conversion families include nitro/electron-withdrawing and "
        "halogenated aryl or heteroaryl aldehydes, alkoxy- and benzyloxy-substituted aryl aldehydes, sulfonyl or "
        "pyridyl aldehydes, and donor-substituted anilino aldehydes. In cobalt, the retained examples are "
        "condition-1 zero-yield alcohol regions. These labels are annotations for interpretation, not mechanistic "
        "assignments, and singleton cobalt clusters are treated as examples rather than as statistical families."
    )


def add_main_results(doc: Document) -> None:
    doc.add_heading("Results and discussion", level=1)
    doc.add_heading("Paired decision-space contrasts reveal boundary undercoverage by 2D acquisition", level=2)
    doc.add_paragraph(
        "The main experiment asks a narrow question: when the same labelled reaction landscape, initial set, "
        "budget and final scoring geometry are held fixed, does changing only the acquisition space alter which "
        "reaction-boundary regions are recovered? The answer is yes. In both main-text systems, the strongest "
        "higher-dimensional or fingerprint-based contrasts increased final-budget boundary coverage relative to "
        "2D acquisition, and the bootstrap intervals remained positive across 20 paired repeats."
    )
    add_picture_if_exists(
        doc,
        CORE_DELTA_FIGURE,
        "Figure 1. Frozen main evidence from the paired r20 contrast. Bars show final-budget boundary-coverage deltas over the matched 2D acquisition baseline; intervals are bootstrap 95% confidence intervals over paired repeats.",
    )
    doc.add_paragraph(
        "The cobalt benchmark gave the largest effect sizes. With weighted iterative CVT acquisition, Tanimoto "
        "space increased boundary coverage by 0.4867 over 2D acquisition, with a 95% confidence interval of "
        "0.4567 to 0.5200. The same dataset also showed positive final-budget gains for CVT in Tanimoto space "
        "and for weighted iterative CVT in full descriptor space. This pattern is important because cobalt is the "
        "small-data case: even when only 60 labelled alcohol entries are available, the projected map and the "
        "decision geometry can lead the campaign toward different zero-yield or near-boundary regions."
    )
    doc.add_paragraph(
        "Aldol provides the broader chemistry-facing support because it contains a larger labelled substrate "
        "landscape with drawable aldehyde structures. In this system, full-space weighted iterative CVT increased "
        "final-budget boundary coverage by 0.1293 relative to the paired 2D baseline, while full-space CVT gave a "
        "similar gain of 0.1183. These effects are smaller than the cobalt effect sizes, but they are more useful "
        "for manuscript interpretation because the missed regions can be tied to recurring aldehyde classes rather "
        "than to isolated examples."
    )

    deltas = pd.read_csv(SPRINT_DIR / "chemical_science_key_boundary_deltas.csv")
    deltas = deltas[deltas["dataset"].isin(["aldol", "cobalt"])].copy()
    switch_columns(doc, 1)
    table = doc.add_table(rows=1, cols=6)
    headers = ["Dataset", "Sampler", "Space vs 2D", "Repeats", "Boundary delta", "95% CI"]
    for idx, header in enumerate(headers):
        set_cell_text(table.cell(0, idx), header, bold=True)
    for _, row in deltas.iterrows():
        cells = table.add_row().cells
        set_cell_text(cells[0], str(row["dataset"]))
        set_cell_text(cells[1], str(row["method"]))
        set_cell_text(cells[2], str(row["high_dimension_vs_2d"]))
        set_cell_text(cells[3], str(int(row["paired_repeats"])))
        set_cell_text(cells[4], f"{float(row['boundary_coverage_delta_mean']):.4f}")
        set_cell_text(
            cells[5],
            f"[{float(row['boundary_coverage_delta_ci_low']):.4f}, {float(row['boundary_coverage_delta_ci_high']):.4f}]",
        )
    format_table(table, [1200, 1700, 1300, 900, 1450, 2810])
    add_caption(doc, "Table 1. Key final-budget boundary-coverage deltas from the r20 paired contrast.")
    switch_columns(doc, 2)

    doc.add_paragraph(
        "These gains should not be read as a universal ranking of decision spaces. In aldol, the same full-space "
        "contrasts that improved boundary coverage reduced failure recall and failure-cluster coverage relative "
        "to the 2D baseline. In cobalt, Tanimoto acquisition gave the strongest boundary-coverage gain but reduced "
        "failure-cluster coverage, whereas full-space weighted iterative CVT gave a smaller boundary gain while "
        "also improving failure-cluster coverage. The result is therefore not simply that higher-dimensional "
        "spaces are better. The stronger conclusion is that boundary coverage, failure recall and failure-family "
        "diversity are separable objectives, and a reaction-scope benchmark should report them separately."
    )

    doc.add_heading("Boundary gains correspond to chemically structured missed regions", level=2)
    doc.add_paragraph(
        "The most important manuscript-facing point is that the extra boundary coverage is not abstract numerical "
        "spread. The missed or late-discovered points are chemically interpretable. In aldol, the low-conversion "
        "regions include alkoxy- and benzyloxy-substituted aryl aldehydes, nitro or otherwise electron-withdrawing "
        "aryl aldehydes, sulfonyl or pyridyl aldehydes, donor-substituted anilino aldehydes, and a broader mixed "
        "heteroaryl, halogenated and trifluoromethyl-containing aldehyde family. These labels are still annotations "
        "rather than mechanistic assignments, but they show that 2D undercoverage can correspond to recognizable "
        "substrate classes rather than random points in a plot."
    )
    doc.add_page_break()
    add_picture_if_exists(
        doc,
        PAIRED_MISS_MAIN_FIGURE,
        "Figure 2. Exact paired-miss examples from the aldol replay. The 2D projection is used only as a location index; the examples are representative row-level checks selected by the full-space replay and missed by the paired 2D replay, while the r20 benchmark in Figure 1 and Table 1 remains the statistical result.",
        width=6.25,
    )
    doc.add_paragraph(
        "The full structure galleries and the larger experiment-package panel are therefore assigned to the "
        "Supporting Information. In the main text, the role of Fig. 2 is deliberately narrower: it verifies that "
        "the numerical boundary-coverage effect can correspond to concrete substrate rows and chemically plausible "
        "nearest-success relationships, not to a decorative projection artifact."
    )
    doc.add_page_break()

    families = pd.read_csv(SPRINT_DIR / "chemical_science_failure_family_table.csv")
    keep = (
        ((families["dataset"] == "aldol") & (families["failure_cluster"].astype(str).isin(["1", "2", "3", "5", "7"])))
        | ((families["dataset"] == "cobalt") & (families["failure_cluster"].astype(str).isin(["1"])))
    )
    shown = families[keep].copy()
    label_override = {
        ("aldol", "1"): "Alkoxy- and benzyloxy-substituted aryl aldehydes",
        ("aldol", "2"): "Nitro/EWG and halogenated aryl or heteroaryl aldehydes",
        ("aldol", "3"): "Sulfonyl-substituted aryl and pyridyl aldehydes",
        ("aldol", "5"): "Dialkylamino/anilino aryl aldehydes",
        ("aldol", "7"): "Mixed heteroaryl, halogenated and CF3 aldehydes",
        ("cobalt", "1"): "Diverse condition-1 zero-yield alcohols",
    }
    switch_columns(doc, 1)
    table = doc.add_table(rows=1, cols=5)
    headers = ["Dataset", "Cluster", "Reviewed family label", "n", "Target range"]
    for idx, header in enumerate(headers):
        set_cell_text(table.cell(0, idx), header, bold=True)
    for _, row in shown.iterrows():
        cluster = str(row["failure_cluster"])
        dataset = str(row["dataset"])
        cells = table.add_row().cells
        set_cell_text(cells[0], dataset)
        set_cell_text(cells[1], cluster)
        set_cell_text(cells[2], label_override.get((dataset, cluster), str(row["manuscript_family_label"])))
        set_cell_text(cells[3], str(int(float(row["cluster_size"]))))
        set_cell_text(cells[4], str(row["target_range"]))
    format_table(table, [1100, 850, 4680, 650, 2080])
    add_caption(doc, "Table 2. Chemistry-reviewed failure-region labels retained for the main text.")
    switch_columns(doc, 2)

    doc.add_paragraph(
        "The cobalt structures are treated more cautiously. The strongest cobalt family is a broad set of "
        "condition-1 zero-yield alcohols, including strained, phenoxy, pyridyl, fluorinated and polar alcohol "
        "examples. Several additional cobalt clusters are singletons, so they are useful for illustrating the "
        "types of failures present in the landscape but are not described as statistically stable mechanistic "
        "families. This distinction keeps the chemistry interpretation aligned with the dataset size, and is why "
        "the full cobalt structure gallery is treated as supporting material rather than a main-text figure."
    )


def add_supporting_results(doc: Document) -> None:
    doc.add_heading("Sensitivity analyses support the boundary-coverage claim but expose objective tradeoffs", level=2)
    doc.add_paragraph(
        "The main paired contrast fixes one boundary definition and one failure-cluster setting. To test whether "
        "the conclusion depends on those choices, we varied the boundary quantile and the number of failure "
        "clusters while preserving shared initial sets and fixed full-space evaluation. Full-space acquisition "
        "remained ahead of 2D acquisition for final-budget boundary coverage across all tested settings in both "
        "aldol and cobalt. The full-vs-2D boundary-coverage delta ranged from 0.0433 to 0.1302 in aldol and from "
        "0.1214 to 0.3167 in cobalt, with the lower confidence-bound values remaining positive."
    )
    doc.add_paragraph(
        "The same sensitivity analysis also gives a useful warning. In cobalt, Tanimoto acquisition improved "
        "failure recall and failure-cluster coverage under the stress-test implementation, but its boundary "
        "coverage was lower than the 2D baseline across the tested boundary quantiles. This is not a failure of "
        "the benchmark; it is exactly why the manuscript separates near-boundary coverage from failure discovery. "
        "A chemist seeking representative failures might choose a different acquisition geometry from a chemist "
        "seeking dense coverage of the compatibility boundary."
    )
    doc.add_paragraph(
        "The detailed robustness ranges are reported in the Supporting Information rather than as a main-text "
        "table. That placement keeps the article focused on the paired r20 decision-space contrast while still "
        "making the boundary-quantile and failure-cluster sensitivity checks available for review."
    )

    doc.add_heading("Drawable Buchwald-Hartwig is retained as an external control rather than a headline result", level=2)
    doc.add_paragraph(
        "A drawable Buchwald-Hartwig control was added to test whether the workflow can preserve reaction-component "
        "structures in an external labelled reaction dataset. It is useful for reproducibility and scope, but it "
        "should not be promoted as the main evidence for the paper. In the focused FPS-style contrast, 2D gave "
        "slightly higher boundary coverage than the 32D, 64D, full and Tanimoto spaces. Conversely, the richer "
        "spaces improved failure-cluster coverage from 0.9125 to 1.0000. This mixed outcome is valuable as a "
        "control because it prevents the paper from reading as a one-sided claim that every richer representation "
        "wins every metric."
    )
    doc.add_paragraph(
        "The representative Buchwald-Hartwig structures, together with the full Aldol and cobalt structure "
        "galleries, are therefore better used as Supporting Information figures. In the main text they function "
        "as scope and workflow controls, not as an additional central claim."
    )


def add_discussion(doc: Document) -> None:
    doc.add_heading("Interpretation: 2D maps remain useful interfaces, but not self-validating decision spaces", level=2)
    doc.add_paragraph(
        "Taken together, the results support a deliberately restrained claim. A 2D reaction-scope map can be an "
        "excellent interface for inspection, discussion and communication, but it does not validate itself as the "
        "space in which acquisition decisions should be made. The paired design shows that changing the acquisition "
        "geometry can change which boundary regions are observed under the same budget. The chemistry review then "
        "shows why this matters: undercovered points can belong to recognizable low-conversion or zero-yield "
        "substrate families, not merely to visually inconvenient parts of a projection."
    )
    doc.add_paragraph(
        "This framing also avoids an overclaim that reviewers would rightly challenge. The data do not show that "
        "full descriptor space, Tanimoto space or any other high-dimensional representation is always superior. "
        "Instead, they show that the choice of decision space changes the acquisition objective that is actually "
        "optimized. Full-space aldol acquisition improves near-boundary coverage but gives lower failure recall "
        "and lower failure-cluster coverage in the main contrast; cobalt Tanimoto acquisition gives the largest "
        "boundary-coverage gain in the main contrast but a less balanced failure-family profile than cobalt "
        "full-space acquisition. These tradeoffs are not secondary details. They are the central reason to evaluate "
        "reaction-scope mapping with multiple boundary-aware endpoints."
    )
    doc.add_paragraph(
        "For practical reaction-scope studies, the implication is a reporting rule rather than a single preferred "
        "algorithm. If a visual map is used to guide substrate selection, the manuscript should state which space "
        "was used for acquisition, which space was used for final scoring, and whether boundary coverage, failure "
        "recall and failure-family diversity agree. When they disagree, the disagreement should be interpreted "
        "chemically rather than averaged away into one aggregate score."
    )

    doc.add_heading("Limitations", level=2)
    doc.add_paragraph(
        "Several limitations define the current scope. First, cobalt is a small-data case with a high final-budget "
        "fraction, so it supports the chemistry story but should not carry budget-generalization alone. Aldol is "
        "therefore the stronger main-text structure-interpretation example, while cobalt is best read as a "
        "small-data stress case showing that projection-dependent acquisition can still matter."
    )
    doc.add_paragraph(
        "Second, the failure-region labels are chemistry-facing annotations, not mechanistic assignments. The "
        "labels were generated from full-space failure clusters and simple structure tags, then curated into "
        "manuscript-readable families. They are sufficient to show that missed regions are chemically structured, "
        "but they should not be used to claim a detailed mechanistic origin without further experimental or "
        "computational analysis."
    )
    doc.add_paragraph(
        "Third, the robustness analysis is a sensitivity stress test rather than a replacement for the main paired "
        "contrast. It used a NumPy-only weighted-iterative-CVT-like implementation to avoid local numerical-library "
        "instability, so the main quantitative claim should continue to rest on the frozen r20 paired contrast. "
        "Finally, the drawable Buchwald-Hartwig control and the broader reference-region or Suzuki-Miyaura scans "
        "belong in the supporting information unless the manuscript is later expanded into a broader benchmark "
        "paper."
    )

    doc.add_heading("Conclusions", level=1)
    doc.add_paragraph(
        "A reaction-scope map can be visually useful while remaining insufficiently validated as an acquisition "
        "space. The frozen core evidence supports a concise benchmark article: compare acquisition spaces under "
        "paired retrospective campaigns, score them in a fixed full-space geometry, and report boundary-aware "
        "metrics together with chemistry-facing missed-region examples. The resulting message is modest but "
        "actionable: visualization can guide human interpretation, but boundary-discovery claims require their own "
        "decision-space validation."
    )


def add_compliance_sections(doc: Document) -> None:
    doc.add_heading("Supporting Information", level=1)
    doc.add_paragraph(
        "The Supporting Information contains the full paired replay tables, exact paired-miss examples, full "
        "Aldol and Cobalt structure galleries, chemistry-reviewed failure-region labels, boundary-quantile and "
        "failure-cluster sensitivity analyses, and the drawable Buchwald-Hartwig external control. These materials "
        "are included as review-facing evidence, while the main text is restricted to the paired r20 "
        "decision-space contrast and compact chemistry interpretation."
    )
    doc.add_heading("Author contributions", level=1)
    doc.add_paragraph(
        "Author contributions will be finalized with the CRediT taxonomy before submission. The current draft "
        "anticipates contributions in conceptualization, data curation, software, formal analysis, visualization, "
        "writing-original draft, writing-review and editing, supervision, and project administration."
    )
    doc.add_heading("Conflicts of interest", level=1)
    doc.add_paragraph("There are no conflicts to declare.")
    doc.add_heading("Data availability", level=1)
    doc.add_paragraph(
        "The analysis scripts, configuration files, intermediate summary tables, figure inputs, and "
        "manuscript-generation scripts will be made available in a public repository before submission. The "
        "original Aldol and Cobalt reaction-scope datasets are from the cited ScopeMap study. The drawable "
        "Buchwald-Hartwig control, derived summary tables, and generated manuscript figures are documented in the "
        "Supporting Information. A public repository URL and archival DOI will be added before submission."
    )
    doc.add_heading("Acknowledgements", level=1)
    doc.add_paragraph(
        "Funding, institutional support, contributor acknowledgements, and any journal-required AI-tool disclosure "
        "will be added before submission. AI-assisted drafting and code-generation support, if disclosed, should "
        "be described as assistance with manuscript preparation, analysis scripting, figure assembly, and editorial "
        "revision; scientific interpretation and final responsibility remain with the authors."
    )


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_core_delta_figure()
    build_main_paired_miss_figure()
    doc = Document()
    setup_document(doc)
    add_title_block(doc)
    add_abstract(doc)
    switch_columns(doc, 2)
    add_introduction(doc)
    doc.add_page_break()
    add_methods(doc)
    add_main_results(doc)
    add_supporting_results(doc)
    add_discussion(doc)
    add_compliance_sections(doc)
    add_references(doc)
    try:
        doc.save(DOCX_OUT)
        return DOCX_OUT
    except PermissionError:
        try:
            doc.save(DOCX_FALLBACK_OUT)
            return DOCX_FALLBACK_OUT
        except PermissionError:
            timestamped = OUT_DIR / f"reaction_scope_boundary_core_manuscript_chemical_science_{datetime.now():%Y%m%d_%H%M%S}.docx"
            doc.save(timestamped)
            return timestamped


if __name__ == "__main__":
    print(build())
