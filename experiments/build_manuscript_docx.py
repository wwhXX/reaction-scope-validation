# -*- coding: utf-8 -*-
"""Build the archived broad DOCX draft from early benchmark outputs.

This builder predates the frozen core-manuscript pruning. It is retained for
supplement/background material, but the default manuscript builder is now
``build_core_manuscript_docx.py``.
"""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

_MPL_CACHE = Path(__file__).resolve().parents[1] / ".matplotlib-cache"
_MPL_CACHE.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(_MPL_CACHE))

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "experiments" / "results" / "manuscript"
SUMMARY_DIR = ROOT / "experiments" / "results" / "summary"
ALDOL_DIR = ROOT / "experiments" / "results" / "aldol_core"
COBALT_DIR = ROOT / "experiments" / "results" / "cobalt_core"
PROSPECTIVE_DIR = ROOT / "experiments" / "results" / "prospective_boundary_simulation"
TRADEOFF_FIGURE = SUMMARY_DIR / "expanded_boundary_tradeoff.png"
METRIC_CORRELATION_FIGURE = SUMMARY_DIR / "metric_tradeoff_correlation_heatmap.png"
METRIC_PARETO_FIGURE = SUMMARY_DIR / "metric_tradeoff_pareto_scatter.png"
BUDGET_TRAJECTORY_FIGURE = SUMMARY_DIR / "budget_trajectory_boundary_coverage.png"
DOCX_OUT = OUT_DIR / "archive_dimension_aware_scope_mapping_manuscript_figure_focused.docx"
HD_ADVANTAGE_FIGURE = OUT_DIR / "hd_advantage_over_2d.png"
DIMENSION_TREND_FIGURE = OUT_DIR / "dimension_trend_rf.png"
ARTICLE_REFERENCE_FIGURE = OUT_DIR / "article_reference_coverage.png"
EXTERNAL_HTE_FIGURE = OUT_DIR / "external_hte_dimension_trend.png"
PROSPECTIVE_FIGURE = OUT_DIR / "prospective_boundary_best.png"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def format_table(table, widths_dxa: list[int]) -> None:
    set_table_width(table, widths_dxa)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
            if row_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            set_cell_margins(cell)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = "Caption"
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(31, 77, 120)
    run.italic = True


def _read_repeated_ci_tables() -> dict[str, pd.DataFrame]:
    return {
        "aldol_reg": pd.read_csv(ALDOL_DIR / "aldol_repeated30_minimal_xgb_ci_regression_summary_ci.csv"),
        "aldol_cls": pd.read_csv(ALDOL_DIR / "aldol_repeated30_minimal_xgb_ci_classification_summary_ci.csv"),
        "cobalt_reg": pd.read_csv(COBALT_DIR / "cobalt_condition1_repeated100_minimal_xgb_ci_regression_summary_ci.csv"),
        "cobalt_cls": pd.read_csv(COBALT_DIR / "cobalt_condition1_repeated100_minimal_xgb_ci_classification_summary_ci.csv"),
    }


def _metric_value(df: pd.DataFrame, model: str, dimension: str, metric: str) -> float:
    row = df[(df["model"] == model) & (df["dimension"].astype(str) == dimension)]
    if row.empty:
        raise ValueError(f"Missing {model} / {dimension} / {metric}")
    return float(row.iloc[0][metric])


def build_dimension_evidence_figures() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    tables = _read_repeated_ci_tables()

    comparisons = [
        (
            "Aldol R2",
            _metric_value(tables["aldol_reg"], "RandomForestRegressor", "2", "r2_mean"),
            _metric_value(tables["aldol_reg"], "RandomForestRegressor", "full", "r2_mean"),
        ),
        (
            "Aldol ROC-AUC",
            _metric_value(tables["aldol_cls"], "RandomForestClassifier", "2", "roc_auc_mean"),
            _metric_value(tables["aldol_cls"], "RandomForestClassifier", "full", "roc_auc_mean"),
        ),
        (
            "Aldol balanced\naccuracy",
            _metric_value(tables["aldol_cls"], "RandomForestClassifier", "2", "balanced_accuracy_mean"),
            _metric_value(tables["aldol_cls"], "RandomForestClassifier", "full", "balanced_accuracy_mean"),
        ),
        (
            "Cobalt R2",
            _metric_value(tables["cobalt_reg"], "RandomForestRegressor", "2", "r2_mean"),
            _metric_value(tables["cobalt_reg"], "RandomForestRegressor", "full", "r2_mean"),
        ),
        (
            "Cobalt ROC-AUC",
            _metric_value(tables["cobalt_cls"], "RandomForestClassifier", "2", "roc_auc_mean"),
            _metric_value(tables["cobalt_cls"], "RandomForestClassifier", "full", "roc_auc_mean"),
        ),
        (
            "Cobalt balanced\naccuracy",
            _metric_value(tables["cobalt_cls"], "RandomForestClassifier", "2", "balanced_accuracy_mean"),
            _metric_value(tables["cobalt_cls"], "RandomForestClassifier", "full", "balanced_accuracy_mean"),
        ),
    ]

    labels = [item[0] for item in comparisons]
    deltas = [item[2] - item[1] for item in comparisons]
    colors = ["#3B73B9", "#3B73B9", "#3B73B9", "#2E8B57", "#2E8B57", "#2E8B57"]

    fig, ax = plt.subplots(figsize=(8.1, 4.4))
    bars = ax.bar(labels, deltas, color=colors, width=0.68)
    ax.axhline(0, color="#444444", linewidth=0.8)
    ax.set_ylabel("Full descriptors minus 2D")
    ax.set_title("High-dimensional descriptors improve repeated-split model evidence")
    ax.grid(axis="y", alpha=0.22, linewidth=0.8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="x", labelsize=9)
    for bar, delta in zip(bars, deltas):
        y = bar.get_height()
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            y + (0.012 if y >= 0 else -0.012),
            f"+{delta:.3f}",
            ha="center",
            va="bottom" if y >= 0 else "top",
            fontsize=9,
        )
    fig.tight_layout()
    fig.savefig(HD_ADVANTAGE_FIGURE, dpi=300)
    plt.close(fig)

    def rf_series(df: pd.DataFrame, model: str, metric: str) -> list[float]:
        return [
            _metric_value(df, model, dim, metric)
            for dim in ["2", "64", "128", "full"]
            if ((df["model"] == model) & (df["dimension"].astype(str) == dim)).any()
        ]

    aldol_dims = ["2D", "64D", "128D", "full"]
    cobalt_dims = ["2D", "8D", "16D", "full"]
    fig, axes = plt.subplots(1, 2, figsize=(8.4, 4.2), sharey=False)
    panels = [
        (
            axes[0],
            "Aldol",
            aldol_dims,
            rf_series(tables["aldol_reg"], "RandomForestRegressor", "r2_mean"),
            rf_series(tables["aldol_cls"], "RandomForestClassifier", "roc_auc_mean"),
        ),
        (
            axes[1],
            "Cobalt",
            cobalt_dims,
            [
                _metric_value(tables["cobalt_reg"], "RandomForestRegressor", dim, "r2_mean")
                for dim in ["2", "8", "16", "full"]
            ],
            [
                _metric_value(tables["cobalt_cls"], "RandomForestClassifier", dim, "roc_auc_mean")
                for dim in ["2", "8", "16", "full"]
            ],
        ),
    ]
    for ax, title, dims, r2_values, auc_values in panels:
        x = range(len(dims))
        ax.plot(x, r2_values, marker="o", linewidth=2.2, color="#3B73B9", label="R2")
        ax.plot(x, auc_values, marker="s", linewidth=2.2, color="#2E8B57", label="ROC-AUC")
        ax.set_xticks(list(x), dims)
        ax.set_title(title)
        ax.grid(axis="y", alpha=0.22, linewidth=0.8)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.set_ylim(min(min(r2_values), 0) - 0.08, max(max(auc_values), 0.9) + 0.04)
        ax.legend(frameon=False, fontsize=9, loc="lower right")
    axes[0].set_ylabel("Repeated-split mean metric")
    fig.suptitle("RandomForest performance improves when decisions use richer representations", y=1.02)
    fig.tight_layout()
    fig.savefig(DIMENSION_TREND_FIGURE, dpi=300, bbox_inches="tight")
    plt.close(fig)


def build_article_reference_figure() -> None:
    data = pd.read_csv(SUMMARY_DIR / "article_reference_best_coverage.csv")
    data = data.sort_values("reference_coverage_q10_mean", ascending=True)
    labels = [f"{row.dataset}\n{row.method}, {row.dimension}" for row in data.itertuples()]
    values = data["reference_coverage_q10_mean"].to_numpy()
    lows = values - data["reference_coverage_q10_ci_low"].to_numpy()
    highs = data["reference_coverage_q10_ci_high"].to_numpy() - values

    fig, ax = plt.subplots(figsize=(8.2, 4.2))
    y = range(len(data))
    ax.barh(y, values, xerr=[lows, highs], color="#3B73B9", alpha=0.88, capsize=3)
    ax.set_yticks(list(y), labels)
    ax.set_xlabel("Reference coverage within the closest 10% of sampled-space distances")
    ax.set_title("Best reference-subset coverage across article-test datasets")
    ax.set_xlim(0, max(0.85, values.max() + 0.08))
    ax.grid(axis="x", alpha=0.22)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    for idx, value in enumerate(values):
        ax.text(value + 0.015, idx, f"{value:.3f}", va="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(ARTICLE_REFERENCE_FIGURE, dpi=300)
    plt.close(fig)


def build_external_hte_figure() -> None:
    reg = pd.read_csv(SUMMARY_DIR / "external_yield_repeated_ci_regression_summary.csv")
    cls = pd.read_csv(SUMMARY_DIR / "external_yield_repeated_ci_classification_summary.csv")
    dim_order = ["2", "16", "32", "full"]
    label_order = ["2D", "16D", "32D", "full"]

    fig, axes = plt.subplots(1, 2, figsize=(8.4, 4.2), sharey=False)
    for ax, dataset, title in [
        (axes[0], "buchwald_hartwig", "Buchwald-Hartwig"),
        (axes[1], "suzuki_miyaura", "Suzuki-Miyaura"),
    ]:
        reg_sub = reg[reg["dataset"] == dataset].set_index("dimension").loc[dim_order]
        cls_sub = cls[cls["dataset"] == dataset].set_index("dimension").loc[dim_order]
        x = range(len(dim_order))
        ax.plot(x, reg_sub["r2_mean"], marker="o", linewidth=2.2, color="#3B73B9", label="R2")
        ax.plot(x, cls_sub["roc_auc_mean"], marker="s", linewidth=2.2, color="#2E8B57", label="ROC-AUC")
        ax.set_xticks(list(x), label_order)
        ax.set_title(title)
        ax.grid(axis="y", alpha=0.22)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.set_ylim(-0.08, 1.02)
        ax.legend(frameon=False, fontsize=9, loc="lower right")
    axes[0].set_ylabel("Repeated-split mean metric")
    fig.suptitle("External HTE yield benchmarks: 2D is not a reliable decision space", y=1.02)
    fig.tight_layout()
    fig.savefig(EXTERNAL_HTE_FIGURE, dpi=300, bbox_inches="tight")
    plt.close(fig)


PROSPECTIVE_TASKS = [
    (
        "Aldol",
        "aldol_prospective_boundary_budget80_repeat30",
        "Strongest near-boundary result; full-space CVT reaches boundary coverage >= 0.1 earliest.",
    ),
    (
        "Cobalt",
        "cobalt_prospective_boundary_budget24_repeat50",
        "Small dataset; use as supportive evidence rather than a standalone claim.",
    ),
    (
        "Buchwald-Hartwig",
        "buchwald_hartwig_prospective_boundary_budget120_repeat20",
        "Best external support for a boundary-aware decision-space consequence.",
    ),
    (
        "Suzuki-Miyaura",
        "suzuki_miyaura_prospective_boundary_budget120_repeat20",
        "Boundary definition is less sharp because many one-hot points become near-boundary.",
    ),
]


def _format_method_space(row: pd.Series) -> str:
    return f"{row['method']} / {row['dimension']}"


def _fastest_threshold(times: pd.DataFrame, metric: str, threshold: float) -> str:
    subset = times[
        (times["metric"] == metric)
        & (times["threshold"] == threshold)
        & (times["reached_rate"] > 0)
        & times["first_budget_mean"].notna()
    ].copy()
    if subset.empty:
        return "Not reached"
    row = subset.sort_values(["first_budget_mean", "reached_rate"], ascending=[True, False]).iloc[0]
    return f"{_format_method_space(row)}, budget {row['first_budget_mean']:.1f}"


def prospective_summary_rows() -> list[dict[str, str | float]]:
    rows: list[dict[str, str | float]] = []
    for dataset, prefix, note in PROSPECTIVE_TASKS:
        summary = pd.read_csv(PROSPECTIVE_DIR / f"{prefix}_summary.csv")
        final_budget = float(summary["budget_spent"].max())
        final = summary[summary["budget_spent"] == final_budget].copy()
        best_boundary = final.sort_values("boundary_coverage_mean", ascending=False).iloc[0]
        best_cluster = final.sort_values("failure_cluster_coverage_mean", ascending=False).iloc[0]
        times = pd.read_csv(PROSPECTIVE_DIR / f"{prefix}_discovery_times_summary.csv")
        rows.append(
            {
                "dataset": dataset,
                "budget": final_budget,
                "best_boundary_method": _format_method_space(best_boundary),
                "boundary_coverage": float(best_boundary["boundary_coverage_mean"]),
                "boundary_enrichment": float(best_boundary["boundary_enrichment_mean"]),
                "best_cluster_method": _format_method_space(best_cluster),
                "cluster_coverage": float(best_cluster["failure_cluster_coverage_mean"]),
                "fast_boundary": _fastest_threshold(times, "boundary_coverage", 0.10),
                "fast_cluster": _fastest_threshold(times, "failure_cluster_coverage", 1.00),
                "note": note,
            }
        )
    return rows


def build_prospective_figure() -> None:
    rows = prospective_summary_rows()
    labels = [f"{row['dataset']}\n{row['best_boundary_method']}" for row in rows]
    values = [float(row["boundary_coverage"]) for row in rows]
    enrichments = [float(row["boundary_enrichment"]) for row in rows]

    fig, ax = plt.subplots(figsize=(8.5, 4.2))
    y = range(len(rows))
    ax.barh(y, values, color="#2E8B57", alpha=0.88)
    ax.set_yticks(list(y), labels)
    ax.set_xlabel("Final-budget near-boundary coverage")
    ax.set_title("Retrospective prospective simulation: best boundary discovery at final budget")
    ax.grid(axis="x", alpha=0.22)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    for idx, (value, enrichment) in enumerate(zip(values, enrichments)):
        ax.text(value + 0.006, idx, f"{value:.3f}, {enrichment:.2f}x", va="center", fontsize=9)
    ax.set_xlim(0, max(values) + 0.08)
    fig.tight_layout()
    fig.savefig(PROSPECTIVE_FIGURE, dpi=300, bbox_inches="tight")
    plt.close(fig)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(89, 89, 89)


def add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(
        "Visualization Space Is Not Validation Space: "
        "A Dimension-Aware Benchmark for AI-Assisted Reaction Scope Mapping"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Manuscript draft for a short benchmark / evaluation protocol article")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(89, 89, 89)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    run = meta.add_run("Target venue under consideration: Digital Discovery | Draft generated from local benchmark outputs")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_abstract(doc: Document) -> None:
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "Two-dimensional reaction-scope maps are useful visual summaries, but they should not be assumed to "
        "be valid algorithmic decision spaces. This short manuscript draft emphasizes the figures and tables "
        "supporting that claim. Across the primary case-study datasets, article-test reference-subset coverage "
        "tasks, and two external public HTE yield benchmarks, low-dimensional representations lose substantial "
        "modeling and sampling signal relative to medium/high-dimensional spaces. The external HTE benchmarks "
        "make the scale of the effect especially clear: RandomForest R2 increases from -0.019 to 0.870 for "
        "Buchwald-Hartwig and from 0.071 to 0.838 for Suzuki-Miyaura when moving from 2D to full one-hot "
        "reaction-component features. Article-test reference coverage is treated as a reference-region recovery "
        "check, not as boundary-discovery evidence. The labelled benchmarks then show the central multi-objective "
        "result: failure recall, failure-cluster coverage, near-boundary coverage, and global coverage do not "
        "rank sampling strategies interchangeably. Finally, "
        "retrospective prospective simulations show the practical consequence of representation and sampler "
        "choice under a fixed experimental budget: full-space CVT reaches Aldol boundary coverage >= 0.1 earliest, "
        "whereas 32D weighted iterative CVT gives the strongest final-budget near-boundary enrichment on the "
        "Buchwald-Hartwig HTE benchmark. The recommended reporting standard is therefore simple: separate "
        "visualization space from decision space, use imbalance-aware metrics, and report boundary/reference "
        "coverage directly."
    )


def add_key_findings(doc: Document) -> None:
    doc.add_heading("Key Findings", level=1)
    findings = [
        "2D is not merely weaker than full space; in several yield benchmarks it is close to unusable as a predictive decision space.",
        "Article-test reference coverage is method- and representation-dependent, but it is a reference-region recovery task rather than boundary-discovery evidence.",
        "Labelled datasets show explicit metric conflicts: methods that maximize near-boundary coverage are often poor under failure-cluster coverage, failure recall, or global coverage.",
        "Multiple Pareto-optimal method/space combinations remain, supporting an objective-first benchmark rather than a universal sampler ranking.",
        "The full-space weighted iterative CVT baseline is strongest on the Alcohols reference-coverage task, but is not universally dominant.",
        "External HTE benchmarks independently reproduce the dimension-sensitivity pattern using one-hot reaction-condition features.",
        "Batch-by-batch retrospective prospective simulations convert the metric story into an experimental-budget trajectory, not only a final-score table.",
        "Buchwald-Hartwig provides the clearest external prospective-simulation support: 32D weighted iterative CVT gives the strongest final-budget near-boundary enrichment.",
        "The manuscript should foreground figures and summary tables; the prose should mainly explain what each comparison means.",
    ]
    for item in findings:
        doc.add_paragraph(item, style="List Bullet")


def add_intro(doc: Document) -> None:
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "The core question is narrow: when a reaction-scope workflow shows a two-dimensional chemical map, "
        "can that same space also support sampling decisions and validation? The results below argue no. "
        "The manuscript is therefore organized around figures and tables rather than a long narrative."
    )


def add_methods(doc: Document) -> None:
    doc.add_heading("2. Methods", level=1)
    doc.add_heading("2.1 Datasets", level=2)
    doc.add_paragraph(
        "The benchmark now includes Aldol, cobalt, four article-test reference-subset tasks, and two external "
        "public HTE yield datasets from rxn4chemistry/rxn_yields. The Suzuki workbook lacks complete component "
        "SMILES, so the external HTE tasks use one-hot reaction-component features rather than invented structures."
    )
    doc.add_heading("2.2 Comparisons", level=2)
    doc.add_paragraph(
        "Modeling tasks report repeated train/test splits with bootstrap confidence intervals. Sampling tasks "
        "report boundary coverage, failure-cluster coverage, and reference-subset coverage. The weighted iterative CVT "
        "reference-sampling baseline is a high-dimensional iterative CVT variant in which already selected samples "
        "act as repulsion anchors."
    )
    doc.add_paragraph(
        "For the new retrospective prospective simulations, the full labelled table is treated as an oracle. "
        "Each method begins from the same initial labelled set, selects new substrates batch by batch, and only "
        "receives labels for the selected batch. Evaluation is retrospective: full labels define failure clusters "
        "and near-boundary points, then the simulated campaign is scored by how quickly it discovers those regions."
    )
    doc.add_heading("2.3 Experimental-Budget Settings", level=2)
    doc.add_paragraph(
        "Simulation budgets are chosen to compare sequential discovery under realistic but not uniquely optimal "
        "assay limits. Aldol uses 80 selected entries, about 7% of its labelled space. Buchwald-Hartwig and "
        "Suzuki-Miyaura use 120 selected entries, about 2-3% of their labelled spaces. Cobalt uses 24 selected "
        "entries because the dataset contains only 60 rows; this 40% fraction is higher than the other cases, "
        "so Cobalt is treated as supportive small-data evidence rather than the main budget claim."
    )
    doc.add_paragraph(
        "The budget trajectory analysis reports intermediate checkpoints from the same retrospective campaigns. "
        "It should be read as a sensitivity check over available assay budgets, whereas a full supplementary "
        "budget-sensitivity experiment would rerun each method under several independent final-budget settings."
    )
    add_note(
        doc,
        "Terminology: the full-space weighted iterative CVT baseline is a benchmark adaptation, not a line-by-line "
        "reproduction of the original interactive workflow."
    )


def add_results(doc: Document) -> None:
    doc.add_heading("3. Results", level=1)
    doc.add_heading("3.1 Two-Dimensional Representations Lose Predictive Information", level=2)
    doc.add_paragraph(
        "In the lightweight Ridge sanity check on Aldol, two-dimensional PCA explained only about 3.7% "
        "of descriptor variance and gave weaker regression performance than 64D, 128D, or full descriptors. "
        "The stronger repeated-split benchmark preserved the same broad pattern: medium- and high-dimensional "
        "representations were more reliable algorithmic decision spaces than 2D embeddings. With RandomForest "
        "models, full descriptors improved Aldol R2 from -0.055 to 0.225 and ROC-AUC from 0.618 to 0.807. "
        "The cobalt benchmark showed the same direction, with R2 improving from 0.087 to 0.403 and ROC-AUC "
        "from 0.722 to 0.858."
    )
    table = doc.add_table(rows=1, cols=4)
    headers = ["Representation", "Explained variance", "RMSE", "R2"]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text
    rows = [
        ("2D PCA", "0.0374", "17.9417", "0.0797"),
        ("64D PCA", "0.3612", "16.8611", "0.1876"),
        ("128D PCA", "0.5348", "16.6377", "0.2082"),
        ("Full descriptors", "1.0000", "16.5876", "0.2114"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    format_table(table, [2600, 2400, 2000, 2000])
    add_caption(doc, "Table 1. Aldol Ridge sanity check: representation dimensionality affects regression signal.")

    if HD_ADVANTAGE_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(HD_ADVANTAGE_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 1. Repeated-split RandomForest evidence for high-dimensional descriptors over 2D. Bars show "
            "the full-descriptor mean metric minus the 2D mean metric."
        )

    if DIMENSION_TREND_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(DIMENSION_TREND_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 2. RandomForest R2 and ROC-AUC across representation dimensionality. The 2D space is weakest "
            "on both case studies, whereas full descriptors give the strongest combined evidence."
        )

    doc.add_heading("3.2 External HTE Benchmarks Reproduce the Same Pattern", level=2)
    if EXTERNAL_HTE_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(EXTERNAL_HTE_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 3. External public HTE yield benchmarks. Even with one-hot reaction-component features, "
            "2D retains too little decision signal compared with 16D, 32D, and full space."
        )

    external_reg = pd.read_csv(SUMMARY_DIR / "external_yield_repeated_ci_regression_summary.csv")
    external_cls = pd.read_csv(SUMMARY_DIR / "external_yield_repeated_ci_classification_summary.csv")
    table = doc.add_table(rows=1, cols=6)
    headers = ["Dataset", "2D R2", "Full R2", "2D AUC", "Full AUC", "Takeaway"]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text
    for dataset, label in [("buchwald_hartwig", "Buchwald-Hartwig"), ("suzuki_miyaura", "Suzuki-Miyaura")]:
        reg = external_reg[external_reg["dataset"] == dataset].set_index("dimension")
        cls = external_cls[external_cls["dataset"] == dataset].set_index("dimension")
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{reg.loc['2', 'r2_mean']:.3f}"
        cells[2].text = f"{reg.loc['full', 'r2_mean']:.3f}"
        cells[3].text = f"{cls.loc['2', 'roc_auc_mean']:.3f}"
        cells[4].text = f"{cls.loc['full', 'roc_auc_mean']:.3f}"
        cells[5].text = "2D is not a reliable prediction space"
    format_table(table, [1700, 1050, 1050, 1050, 1050, 3100])
    add_caption(doc, "Table 2. External HTE repeated-CI summary, RandomForest models, yield >= 50 for classification.")

    doc.add_heading("3.3 Article-Test Reference Coverage", level=2)
    doc.add_paragraph(
        "The article-test datasets are intentionally interpreted more narrowly than the labelled HTE datasets. "
        "They contain a candidate space plus a known experimental/reference subset, but they do not provide "
        "full-space success/failure labels. Therefore Figure 4 is a reference-region recovery check, not evidence "
        "that any method discovers reaction boundaries in these article-test spaces."
    )
    if ARTICLE_REFERENCE_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ARTICLE_REFERENCE_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 4. Best reference-subset coverage across article-test datasets. Coverage is measured against "
            "known experimental/reference subsets without treating untested molecules as negative labels; it should "
            "not be interpreted as success/failure boundary discovery."
        )

    article_best = pd.read_csv(SUMMARY_DIR / "article_reference_best_coverage.csv")
    table = doc.add_table(rows=1, cols=6)
    headers = ["Dataset", "Best method", "Space", "q10 coverage", "q10 enrich.", "Comment"]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text
    comments = {
        "styrene": "Ward 128D is strongest.",
        "thiol": "Kennard-Stone full wins.",
        "alcohol": "Full-space weighted iterative CVT wins.",
        "oxo_carboxide": "Ward full wins.",
    }
    for _, row in article_best.iterrows():
        cells = table.add_row().cells
        dataset = str(row["dataset"])
        cells[0].text = dataset
        cells[1].text = str(row["method"])
        cells[2].text = str(row["dimension"])
        cells[3].text = f"{row['reference_coverage_q10_mean']:.3f}"
        cells[4].text = f"{row['reference_enrichment_q10_mean']:.2f}x"
        cells[5].text = comments.get(dataset, "")
    format_table(table, [1350, 1700, 900, 1250, 1150, 2870])
    add_caption(doc, "Table 3. Article-test reference coverage, repeat=20.")

    doc.add_heading("3.4 F1 Is Inflated by Aldol Class Imbalance", level=2)
    doc.add_paragraph(
        "Using conversion >= 70 as a success threshold gives a positive-class ratio of approximately 0.865. "
        "A dummy classifier that always predicts success can therefore reach F1 around 0.927 despite having "
        "no boundary-learning ability and ROC-AUC of 0.500. This motivates reporting balanced accuracy, MCC, "
        "AUC, specificity, and PR-AUC alongside F1."
    )

    doc.add_heading("3.5 Expanded Boundary Sampling Results", level=2)
    best_boundary = pd.read_csv(SUMMARY_DIR / "expanded_best_boundary_coverage.csv")
    table = doc.add_table(rows=1, cols=7)
    headers = ["Dataset", "Budget", "Repeats", "Best method", "Space", "Boundary cov.", "Boundary enrich."]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text
    for _, row in best_boundary.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["dataset"])
        cells[1].text = str(int(row["budget"]))
        cells[2].text = str(int(row["repeats"]))
        cells[3].text = str(row["method"])
        cells[4].text = str(row["dimension"])
        ci_low = row.get("boundary_coverage_ci_low")
        ci_high = row.get("boundary_coverage_ci_high")
        if pd.notna(ci_low) and pd.notna(ci_high):
            cells[5].text = f"{row['boundary_coverage_mean']:.4f} [{ci_low:.4f}, {ci_high:.4f}]"
        else:
            cells[5].text = f"{row['boundary_coverage_mean']:.4f}"
        cells[6].text = f"{row['boundary_enrichment_mean']:.4f}"
    format_table(table, [1500, 760, 820, 1700, 980, 2140, 1220])
    add_caption(doc, "Table 4. Best near-boundary coverage in the expanded sampling benchmark.")

    doc.add_paragraph(
        "The expanded benchmark identifies weighted iterative CVT as the strongest near-boundary coverage "
        "method in both case studies. For Aldol, the best combination was weighted iterative CVT in full "
        "descriptor space. For Cobalt, weighted iterative CVT in Tanimoto space was best at thresholds 1, "
        "30, and 50."
    )

    doc.add_heading("3.6 Retrospective Prospective Simulations Show Experimental-Budget Consequences", level=2)
    doc.add_paragraph(
        "The static boundary benchmark asks which method covers boundary points under a fixed sample. The "
        "retrospective prospective simulation asks a more experiment-like question: if a chemist spends a fixed "
        "number of assays in batches, which decision rule discovers boundary regions or failure modes earlier? "
        "This does not replace wet-lab validation, but it is a stronger computational proxy for experimental "
        "consequence than a single final-score benchmark."
    )
    if PROSPECTIVE_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(PROSPECTIVE_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 5. Retrospective prospective simulations. Bars show the best final-budget near-boundary "
            "coverage for each labelled dataset; labels report the winning method/space and enrichment."
        )

    if BUDGET_TRAJECTORY_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(BUDGET_TRAJECTORY_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 6. Budget trajectories from retrospective prospective simulations. Each point shows the "
            "best near-boundary coverage available at that intermediate simulated assay budget."
        )

    prospective_rows = prospective_summary_rows()
    table = doc.add_table(rows=1, cols=6)
    headers = [
        "Dataset",
        "Budget",
        "Best near-boundary campaign",
        "Boundary cov. / enrich.",
        "Fastest full failure-cluster discovery",
        "Interpretation",
    ]
    for idx, text in enumerate(headers):
        table.cell(0, idx).text = text
    for row in prospective_rows:
        cells = table.add_row().cells
        cells[0].text = str(row["dataset"])
        cells[1].text = str(int(float(row["budget"])))
        cells[2].text = str(row["best_boundary_method"])
        cells[3].text = f"{float(row['boundary_coverage']):.4f} / {float(row['boundary_enrichment']):.2f}x"
        cells[4].text = str(row["fast_cluster"])
        cells[5].text = str(row["note"])
    format_table(table, [1150, 740, 1820, 1420, 1880, 2210])
    add_caption(
        doc,
        "Table 5. Batch-by-batch retrospective prospective simulations. Boundary coverage is evaluated at the "
        "final simulated budget; discovery budget reports the first budget at which full failure-cluster coverage "
        "is reached."
    )
    budget_path = SUMMARY_DIR / "budget_trajectory_summary.csv"
    if budget_path.exists():
        budget = pd.read_csv(budget_path)
        table = doc.add_table(rows=1, cols=5)
        headers = ["Dataset", "Final budget", "Best mid-budget", "Best final-budget", "Budget note"]
        for idx, text in enumerate(headers):
            table.cell(0, idx).text = text
        for _, row in budget.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row["dataset"])
            cells[1].text = str(int(float(row["final_budget"])))
            cells[2].text = f"{row['best_mid_boundary']} ({row['best_mid_boundary_coverage']:.3f})"
            cells[3].text = f"{row['best_final_boundary']} ({row['best_final_boundary_coverage']:.3f})"
            first = row["first_budget_boundary_coverage_ge_0_10"]
            if pd.notna(first):
                cells[4].text = f"Boundary coverage >= 0.10 by budget {float(first):.0f}."
            else:
                cells[4].text = "Boundary coverage remains below 0.10 in this budget range."
        format_table(table, [1150, 900, 2250, 2250, 2450])
        add_caption(doc, "Table 6. Budget trajectory check across intermediate simulated assay budgets.")
    doc.add_paragraph(
        "Aldol and Buchwald-Hartwig provide the most useful manuscript signal. On Aldol, full-space CVT gives "
        "the best final-budget boundary coverage and reaches boundary coverage >= 0.1 earliest. On the external "
        "Buchwald-Hartwig HTE benchmark, 32D weighted iterative CVT gives the strongest final-budget boundary "
        "coverage and enrichment. Cobalt is supportive but too small to carry the claim alone. Suzuki-Miyaura "
        "shows that the current nearest-opposite-class boundary definition can become less discriminating for "
        "duplicate-rich one-hot spaces, so it should be presented as a caveat and a target for boundary-definition "
        "sensitivity analysis."
    )

    doc.add_heading("3.7 Cross-Metric Trade-Offs Require Multi-Objective Evaluation", level=2)
    doc.add_paragraph(
        "Figure 4 alone cannot support a multi-objective claim because it evaluates only reference-region "
        "recovery. The multi-objective evidence must come from labelled datasets where failure recall, "
        "failure-cluster coverage, near-boundary coverage, boundary enrichment, and global coverage can be "
        "computed for the same method/space combinations."
    )
    if METRIC_CORRELATION_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(METRIC_CORRELATION_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 7. Spearman correlations among labelled boundary metrics. Boundary coverage is negatively "
            "correlated with failure-cluster coverage in the static Aldol and Cobalt screens, showing that "
            "these objectives cannot be treated as interchangeable."
        )
    if METRIC_PARETO_FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(METRIC_PARETO_FIGURE), width=Inches(6.3))
        add_caption(
            doc,
            "Figure 8. Pareto-frontier view of the static boundary benchmark. Larger outlined points are "
            "Pareto-optimal under failure recall, failure-cluster coverage, near-boundary coverage, boundary "
            "enrichment, and coverage compactness."
        )
    conflict_path = SUMMARY_DIR / "metric_tradeoff_rank_conflicts.csv"
    if conflict_path.exists():
        conflicts = pd.read_csv(conflict_path)
        conflicts = conflicts[
            (conflicts["primary_metric"] == "boundary_coverage_mean")
            & (conflicts["secondary_metric"] == "failure_cluster_coverage_mean")
            & (conflicts["analysis_type"].isin(["static_boundary", "prospective_final"]))
        ].copy()
        preferred_order = [
            "aldol",
            "cobalt_condition1",
            "cobalt_threshold50",
            "buchwald_hartwig",
            "suzuki_miyaura",
        ]
        conflicts["dataset_order"] = conflicts["dataset"].map(
            {dataset: idx for idx, dataset in enumerate(preferred_order)}
        ).fillna(99)
        conflicts = conflicts.sort_values(["analysis_type", "dataset_order"]).head(7)
        table = doc.add_table(rows=1, cols=6)
        headers = [
            "Analysis",
            "Dataset",
            "Best boundary",
            "Rank on clusters",
            "Best clusters",
            "Rank on boundary",
        ]
        for idx, text in enumerate(headers):
            table.cell(0, idx).text = text
        for _, row in conflicts.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row["analysis_type"])
            cells[1].text = str(row["dataset"])
            cells[2].text = str(row["primary_best"])
            cells[3].text = str(int(row["primary_best_rank_on_secondary"]))
            cells[4].text = str(row["secondary_best"])
            cells[5].text = str(int(row["secondary_best_rank_on_primary"]))
        format_table(table, [1300, 1500, 1900, 1100, 1900, 1100])
        add_caption(
            doc,
            "Table 7. Rank conflicts between near-boundary coverage and failure-cluster coverage. A high rank "
            "number means that the metric winner performs poorly when judged by the other objective."
        )
    doc.add_paragraph(
        "The trade-off analysis changes the manuscript logic. The claim is not that one sampler dominates all "
        "others; instead, the benchmark shows that the preferred sampler depends on the experimental objective. "
        "For example, the best near-boundary method can rank near the bottom for failure-cluster coverage, while "
        "the best failure-cluster method can rank poorly for near-boundary coverage. Multiple Pareto-optimal "
        "method/space combinations remain, which supports an objective-first workflow: define the experimental "
        "objective, choose the metric, then compare samplers under that metric."
    )


def add_discussion(doc: Document) -> None:
    doc.add_heading("4. Discussion", level=1)
    doc.add_paragraph(
        "The figures support a compact argument. Two-dimensional maps are useful for inspection, but the "
        "decision-space penalty is large enough that they should not be used for sampling or validation by default. "
        "The HTE one-hot benchmarks are especially useful because they show the same failure mode outside the "
        "the two primary molecular-fingerprint case-study datasets."
    )
    doc.add_paragraph(
        "The method lesson is also practical: no single sampling score captures reference coverage, boundary "
        "coverage, failure recall, failure-cluster discovery, and global coverage. The manuscript should therefore "
        "present a figure-led evaluation protocol rather than a claim that any one sampler is universally best."
    )
    doc.add_paragraph(
        "The retrospective prospective simulations make the argument more concrete. They turn a representation "
        "choice into an experimental-budget consequence: which simulated campaign would reach boundary or failure "
        "coverage sooner if labels were revealed only after each selected batch? The strongest version of this "
        "evidence is currently Aldol plus Buchwald-Hartwig; Cobalt is small, and Suzuki-Miyaura exposes a useful "
        "limitation of the present boundary definition in duplicate-rich one-hot feature spaces."
    )
    doc.add_paragraph(
        "The cross-metric analysis also clarifies the role of the article-test datasets. Those datasets are useful "
        "for reference-region recovery because they resemble published substrate-scope settings, but they should "
        "not carry the main boundary-discovery or trade-off claim. The labelled datasets are the right evidence "
        "base for metric conflict, Pareto frontiers, and budget trajectories."
    )


def add_limitations(doc: Document) -> None:
    doc.add_heading("5. Limitations", level=1)
    limitations = [
        "The benchmark now spans published case-study data, article-test reference subsets, and two external HTE yield datasets, but the external tasks use reaction-component one-hot features rather than molecular descriptors.",
        "The weighted iterative CVT implementation is a benchmark-oriented baseline adaptation, not a full reproduction of the original interactive workflow.",
        "Some deterministic samplers yield zero-width confidence intervals under fixed benchmark settings; these should be interpreted as deterministic algorithm behavior, not as biological or chemical certainty.",
        "Boundary definitions depend on the chosen representation and nearest-opposite-class distance threshold; Suzuki-Miyaura shows that this definition can become less discriminating in duplicate-rich one-hot spaces.",
        "Budget trajectories are reported from intermediate checkpoints of the same retrospective prospective campaigns; a full supplementary budget-sensitivity study should rerun each method under several independent final-budget settings.",
        "Article-test reference-subset results should be interpreted as reference-region recovery only, because untested candidates are not reliable negative labels.",
        "The retrospective prospective simulations are stronger than static retrospective scores, but they still do not replace new wet-lab validation.",
    ]
    for item in limitations:
        doc.add_paragraph(item, style="List Bullet")


def add_conclusion(doc: Document) -> None:
    doc.add_heading("6. Conclusions", level=1)
    doc.add_paragraph(
        "The manuscript should make one disciplined claim: visualization space is not validation space. The "
        "current figures show that 2D can lose enough signal to become unsuitable for modeling, reference coverage, "
        "or boundary evaluation. The new batch-by-batch simulations add an experimental-budget interpretation: "
        "decision-space and sampler choices can change how quickly a simulated campaign discovers boundary regions "
        "and failure modes. The cross-metric analysis adds the missing logic for multi-objective evaluation: "
        "reference recovery, failure recall, failure-cluster coverage, near-boundary coverage, and global coverage "
        "can favor different method/space choices. A stronger reporting standard should separate visualization, "
        "decision, prediction, and sampling-evaluation spaces."
    )


def add_availability_and_refs(doc: Document) -> None:
    doc.add_heading("Data and Code Availability", level=1)
    doc.add_paragraph(
        "Core benchmark outputs are generated from experiments/run_benchmark.py, and cross-metric trade-off outputs "
        "are generated from experiments/analyze_metric_tradeoffs.py. Key configs are article_reference_sampling.json, "
        "external_yield_benchmarks.json, prospective_boundary_simulation.json, boundary_expanded.json, and "
        "cobalt_threshold_expanded_boundary.json. "
        "Summary outputs are located under experiments/results/summary."
    )
    doc.add_heading("References", level=1)
    refs = [
        "Li et al. ScopeMap: an AI-assisted human-in-the-loop workflow for mapping reaction scope. Angewandte Chemie, 2026. Details to be completed from the final citation record.",
        "Standard references for Morgan/ECFP fingerprints, MACCS keys, PCA, Tanimoto distance, CVT sampling, Kennard-Stone sampling, and class-imbalance metrics should be added during manuscript polishing.",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="List Number")


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run("Draft manuscript - decision-space validation benchmark")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(117, 117, 117)


def build() -> None:
    print(
        "Archived broad builder: use experiments/build_core_manuscript_docx.py "
        "for the frozen Aldol/Cobalt core manuscript."
    )
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_dimension_evidence_figures()
    build_article_reference_figure()
    build_external_hte_figure()
    build_prospective_figure()
    doc = Document()
    configure_styles(doc)
    add_title_block(doc)
    add_abstract(doc)
    add_key_findings(doc)
    add_intro(doc)
    add_methods(doc)
    add_results(doc)
    add_discussion(doc)
    add_limitations(doc)
    add_conclusion(doc)
    add_availability_and_refs(doc)
    add_footer(doc)
    try:
        doc.save(DOCX_OUT)
        saved_path = DOCX_OUT
    except PermissionError:
        saved_path = DOCX_OUT.with_name(f"{DOCX_OUT.stem}_updated_{datetime.now():%Y%m%d_%H%M%S}{DOCX_OUT.suffix}")
        doc.save(saved_path)
    print(saved_path)


if __name__ == "__main__":
    build()
