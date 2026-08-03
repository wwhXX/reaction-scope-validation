from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_DIR = ROOT / "experiments" / "results" / "manuscript"
SOURCE = MANUSCRIPT_DIR / "reaction_scope_boundary_core_manuscript_chemical_science_with_candidate_refs.docx"
OUTPUT = MANUSCRIPT_DIR / "reaction_scope_boundary_core_manuscript_chemical_science_revised_20260802.docx"
DETAIL = ROOT / "experiments" / "results" / "chemical_science_sprint" / "decision_space_contrast_r20_detail.csv"
FIG1 = MANUSCRIPT_DIR / "boundary_coverage_absolute_revised_20260802.png"
FIG2_SOURCE = MANUSCRIPT_DIR / "scopehd_vs_2d_boundary_comparison_main.png"
FIG2 = MANUSCRIPT_DIR / "scopehd_exact_miss_examples_revised_20260802.png"


def bootstrap_ci(values: np.ndarray, samples: int = 2000, seed: int = 42) -> tuple[float, float]:
    rng = np.random.default_rng(seed)
    values = np.asarray(values, dtype=float)
    idx = rng.integers(0, len(values), size=(samples, len(values)))
    means = values[idx].mean(axis=1)
    return float(np.quantile(means, 0.025)), float(np.quantile(means, 0.975))


def build_figure_1() -> None:
    detail = pd.read_csv(DETAIL)
    detail = detail.loc[detail["budget_spent"].eq(detail.groupby("dataset")["budget_spent"].transform("max"))]
    contrasts = [
        ("Aldol", "aldol", "weighted_itr_cvt", "full", "Full descriptor"),
        ("Cobalt", "cobalt", "weighted_itr_cvt", "tanimoto", "Tanimoto"),
    ]
    records: list[dict[str, float | str]] = []
    for panel, dataset, method, richer, richer_label in contrasts:
        group = detail.loc[(detail["dataset"] == dataset) & (detail["method"] == method)]
        for dim, label in [("2", "2D"), (richer, richer_label)]:
            values = group.loc[group["decision_dimension"].astype(str) == dim, "boundary_coverage"].to_numpy(float)
            lo, hi = bootstrap_ci(values, seed=42 + len(records) * 101)
            records.append({"panel": panel, "space": label, "mean": values.mean(), "lo": lo, "hi": hi})

    canvas = Image.new("RGB", (2400, 860), "white")
    draw = ImageDraw.Draw(canvas)
    regular_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    title_font = ImageFont.truetype(bold_path, 52)
    panel_font = ImageFont.truetype(bold_path, 43)
    label_font = ImageFont.truetype(regular_path, 34)
    value_font = ImageFont.truetype(bold_path, 34)
    small_font = ImageFont.truetype(regular_path, 30)

    title = "Paired decision-space audit with shared initial sets"
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((2400 - (title_box[2] - title_box[0])) / 2, 25), title, fill="#17324D", font=title_font)
    colors = ["#4C78A8", "#2E8B57"]
    max_value = 0.75
    plot_top, plot_bottom = 160, 680
    for panel_i, panel in enumerate(["Aldol", "Cobalt"]):
        left = 170 + panel_i * 1150
        right = left + 930
        rows = [row for row in records if row["panel"] == panel]
        panel_box = draw.textbbox((0, 0), panel, font=panel_font)
        draw.text(((left + right - (panel_box[2] - panel_box[0])) / 2, 92), panel, fill="#17324D", font=panel_font)
        for tick in [0.0, 0.25, 0.5, 0.75]:
            y = plot_bottom - int((tick / max_value) * (plot_bottom - plot_top))
            draw.line((left, y, right, y), fill="#D9DEE5", width=2)
            if panel_i == 0:
                draw.text((left - 82, y - 17), f"{tick:.2f}", fill="#555555", font=small_font)
        draw.line((left, plot_top, left, plot_bottom), fill="#555555", width=3)
        draw.line((left, plot_bottom, right, plot_bottom), fill="#555555", width=3)
        centers = [left + 290, left + 670]
        for row_i, (row, x_center) in enumerate(zip(rows, centers)):
            mean = float(row["mean"])
            lo = float(row["lo"])
            hi = float(row["hi"])
            bar_width = 210
            y_mean = plot_bottom - int((mean / max_value) * (plot_bottom - plot_top))
            y_lo = plot_bottom - int((lo / max_value) * (plot_bottom - plot_top))
            y_hi = plot_bottom - int((hi / max_value) * (plot_bottom - plot_top))
            draw.rectangle((x_center - bar_width // 2, y_mean, x_center + bar_width // 2, plot_bottom), fill=colors[row_i])
            draw.line((x_center, y_hi, x_center, y_lo), fill="#222222", width=5)
            draw.line((x_center - 30, y_hi, x_center + 30, y_hi), fill="#222222", width=5)
            draw.line((x_center - 30, y_lo, x_center + 30, y_lo), fill="#222222", width=5)
            value = f"{mean:.3f}"
            value_box = draw.textbbox((0, 0), value, font=value_font)
            draw.text((x_center - (value_box[2] - value_box[0]) / 2, y_hi - 55), value, fill="#111111", font=value_font)
            label = str(row["space"])
            label_box = draw.textbbox((0, 0), label, font=label_font)
            draw.text((x_center - (label_box[2] - label_box[0]) / 2, plot_bottom + 25), label, fill="#222222", font=label_font)
    y_label = "Final-budget boundary coverage"
    label_layer = Image.new("RGBA", (700, 60), (255, 255, 255, 0))
    label_draw = ImageDraw.Draw(label_layer)
    label_draw.text((0, 0), y_label, fill="#222222", font=label_font)
    label_layer = label_layer.rotate(90, expand=True)
    canvas.paste(label_layer, (15, 180), label_layer)
    canvas.save(FIG1, dpi=(300, 300))


def build_figure_2() -> None:
    image = Image.open(FIG2_SOURCE).convert("RGB")
    width, height = image.size
    # Retain the three chemistry-facing row cards; titles and the projection index move to the caption/SI.
    cropped = image.crop((0, 245, int(width * 0.52), height - 50))
    cropped.save(FIG2, quality=95)


def find_paragraph(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise KeyError(prefix)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    cursor = 0
    for match in re.finditer(r"\[\[CITE:([^\]]+)\]\]", text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        run = paragraph.add_run(match.group(1))
        run.font.superscript = True
        run.font.size = Pt(7)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def set_caption(paragraph, label: str, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lead = paragraph.add_run(label)
    lead.bold = True
    lead.italic = False
    lead.font.color.rgb = RGBColor(0, 0, 0)
    lead.font.size = Pt(9)
    rest = paragraph.add_run(text)
    rest.bold = False
    rest.italic = False
    rest.font.color.rgb = RGBColor(0, 0, 0)
    rest.font.size = Pt(9)


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def set_section_continuous(paragraph) -> None:
    sect_pr = paragraph._p.get_or_add_pPr().sectPr
    if sect_pr is None:
        raise RuntimeError("Expected a section break")
    type_el = sect_pr.find(qn("w:type"))
    if type_el is None:
        type_el = OxmlElement("w:type")
        sect_pr.insert(0, type_el)
    type_el.set(qn("w:val"), "continuous")


def replace_picture(paragraph, path: Path, width: float) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_i, row in enumerate(table.rows):
        for col_i, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.3)
                    run.font.name = "Times New Roman"
                    if row_i == 0:
                        run.bold = True


def rebuild_result_table(table) -> None:
    headers = ["Dataset", "Sampler", "Richer space", "Coverage\n2D to richer", "Delta", "95% CI"]
    rows = [
        ["cobalt", "weighted itr. CVT", "Tanimoto", "0.1633 -> 0.6500", "0.4867", "[0.4567, 0.5200]"],
        ["cobalt", "CVT", "Tanimoto", "0.3667 -> 0.6500", "0.2833", "[0.2600, 0.3100]"],
        ["cobalt", "weighted itr. CVT", "full", "0.1633 -> 0.4233", "0.2600", "[0.2300, 0.2900]"],
        ["aldol", "weighted itr. CVT", "full", "0.0568 -> 0.1861", "0.1293", "[0.1236, 0.1355]"],
        ["aldol", "CVT", "full", "0.0725 -> 0.1908", "0.1183", "[0.1141, 0.1229]"],
    ]
    for row_i, values in enumerate([headers, *rows]):
        for col_i, value in enumerate(values):
            table.cell(row_i, col_i).text = value
    style_table(table, [0.6, 1.35, 0.9, 1.4, 0.65, 1.35])


def revise() -> None:
    build_figure_1()
    build_figure_2()
    doc = Document(SOURCE)

    # The first two-column section previously started on a new page, leaving page 2 almost empty.
    assert doc.paragraphs[6].text == ""
    set_section_continuous(doc.paragraphs[6])

    set_paragraph(
        find_paragraph(doc, "Reaction-scope maps are increasingly used"),
        "Reaction-scope maps increasingly guide substrate selection in human-in-the-loop and data-driven workflows. "
        "The two-dimensional display used to discuss a scope, however, need not preserve the geometry required for acquisition. "
        "We tested this distinction in paired retrospective campaigns on the ScopeMap aldol and cobalt landscapes. "
        "Within each repeat, the labelled dataset, initial set, budget, sampler and full-space evaluation geometry were fixed, while acquisition was performed in either a 2D projection or a richer descriptor or Tanimoto space. "
        "For weighted iterative CVT, final-budget boundary coverage increased from 0.163 to 0.650 in cobalt when acquisition moved from 2D to Tanimoto space, and from 0.057 to 0.186 in aldol when acquisition moved from 2D to the full descriptor space. "
        "The corresponding paired gains were 0.4867 and 0.1293, with positive bootstrap intervals. "
        "The additional aldol points included recurring nitro/electron-withdrawing, alkoxy/benzyloxy and sulfonyl/pyridyl aldehyde regions, although not every failure cluster formed a homogeneous chemical family. "
        "Sensitivity analyses and an external Buchwald-Hartwig control further showed that boundary coverage and failure-family discovery can favour different decision spaces. "
        "Thus, 2D maps remain useful interfaces, but acquisition claims require explicit validation in the decision space used to select experiments.",
    )
    set_paragraph(
        find_paragraph(doc, "Here we treat this as an empirical benchmark problem"),
        "Here we test whether acquisition geometry changes boundary recovery under a fixed budget. Using the published ScopeMap aldol and cobalt landscapes, we hold the initial set, sampler, budget and full-space scoring geometry constant while varying acquisition between 2D and richer spaces. "
        "We report boundary coverage, failure recall, failure-cluster coverage and boundary enrichment, and relate selected differences to representative substrate subsets. Our claim is limited to a reporting principle: visual maps are useful interfaces, but boundary-discovery claims require independent decision-space validation [[CITE:1,12-18,22-24]].",
    )

    set_paragraph(
        find_paragraph(doc, "The benchmark was designed as a paired retrospective"),
        "We used a paired retrospective acquisition design to isolate the effect of decision-space geometry. "
        "For each repeat, the labelled dataset, binary target, initial selected set, final budget, sampler and evaluation space were held constant. "
        "Only the representation used to select subsequent batches was changed. Full labels were available to the retrospective oracle but were revealed to the acquisition policy only after a simulated batch had been selected. "
        "All boundary metrics were calculated in one fixed full-descriptor evaluation geometry, so differences between paired campaigns reflect acquisition rather than a change in the scoring space [[CITE:1,12-18,22-28,46-48]].",
    )

    # Remove project-management language and its evidence-hierarchy table.
    frozen_heading = find_paragraph(doc, "Frozen evidence hierarchy")
    frozen_body = find_paragraph(doc, "The frozen evidence hierarchy")
    table_to_remove = next(table for table in doc.tables if table.cell(0, 0).text.strip() == "Layer")
    paragraph_snapshot = list(doc.paragraphs)
    frozen_index = next(i for i, paragraph in enumerate(paragraph_snapshot) if paragraph._p is frozen_body._p)
    for paragraph in paragraph_snapshot[frozen_index + 1 : frozen_index + 3]:
        if paragraph.text == "":
            remove_paragraph(paragraph)
    table_to_remove._element.getparent().remove(table_to_remove._element)
    remove_paragraph(frozen_body)
    remove_paragraph(frozen_heading)

    set_paragraph(
        find_paragraph(doc, "Each reaction landscape was represented"),
        "Numeric feature columns were extracted after removing identifiers and response columns. For Euclidean decision spaces, features were standardized to zero mean and unit variance. "
        "The 2D spaces were the first two principal components of the standardized matrix; the full spaces retained all standardized features. "
        "For Tanimoto acquisition, the original non-standardized feature matrix was used with generalized Tanimoto distance, defined as one minus the dot-product similarity normalized by the two squared vector norms and their shared dot product. "
        "The main paired contrasts used CVT and weighted iterative CVT as geometric probes of decision-space choice rather than as mechanistic models of reactivity [[CITE:12-18,22-24,32,40]].",
    )
    set_paragraph(
        find_paragraph(doc, "For each repeat, the same initial selected set"),
        "Each repeat began from the same CVT-selected initial set across all decision spaces, generated in the fixed full-space evaluation geometry with base seed 42 and repeat-specific deterministic seeds. "
        "Aldol campaigns started with 10 entries, added batches of 10 and stopped at 80 selected entries; cobalt campaigns started with 6 entries, added batches of 6 and stopped at 24. "
        "CVT selected cluster representatives from the remaining candidates. Weighted iterative CVT initialized batch centres by CVT, retained all previously selected entries as fixed centres, and applied repulsion from previously observed failures before mapping the optimized centres to unique remaining candidates. "
        "The weighted update used 20 iterations, 10 CVT initialization iterations, a learning rate of 0.01 and the implementation's default repulsion strength of 1.0.",
    )
    set_paragraph(
        find_paragraph(doc, "The primary endpoint is final-budget boundary coverage"),
        "Boundary status was defined once in the fixed full-space evaluation geometry. For every labelled entry, we calculated the distance to its nearest member of the opposite outcome class; entries at or below the 25th percentile of this nearest-opposite distance were designated near-boundary points. "
        "Final-budget boundary coverage is the fraction of this fixed boundary set selected by a campaign. Boundary enrichment is the boundary fraction in the selected set divided by the boundary fraction in the full dataset. "
        "Failure recall is the fraction of all failures selected. Failure-cluster coverage is the fraction of full-space failure clusters represented by at least one selected entry; failures were partitioned by k-means into eight clusters for aldol and four for cobalt.",
    )
    set_paragraph(
        find_paragraph(doc, "The paper does not assume that a single decision space"),
        "These endpoints were retained separately because they encode different experimental objectives. A campaign can improve local coverage of the compatibility boundary while recovering fewer failures or fewer failure clusters. "
        "No aggregate score was used to average away these disagreements [[CITE:41,45]].",
    )
    set_paragraph(
        find_paragraph(doc, "The main results use 20 paired repeats"),
        "The primary summaries used 20 paired repeats for each decision-space contrast. For every sampler and dataset, each richer-space campaign was matched to its 2D campaign by repeat and shared initial set. "
        "We report the mean paired difference at the final budget. Percentile 95% confidence intervals were obtained from 2000 bootstrap resamples of the 20 paired differences, using deterministic metric-specific seeds [[CITE:27]]. "
        "Absolute coverage values are reported together with paired differences to retain the scale of each comparison.",
    )
    set_paragraph(
        find_paragraph(doc, "Robustness checks varied the boundary quantile"),
        "Sensitivity analyses varied the boundary quantile and failure-cluster count while preserving shared initial sets and fixed full-space evaluation. "
        "These analyses used a separate NumPy-only weighted-iterative-CVT-like approximation because the primary implementation was unstable in the local numerical environment. "
        "They are therefore used only to assess directional sensitivity to endpoint definitions; their effect sizes are not pooled with, or treated as replacements for, the primary paired estimates.",
    )
    set_paragraph(
        find_paragraph(doc, "Failure-region labels were assigned"),
        "Failure-region labels were assigned after the numerical analysis to support chemical interpretation. "
        "Aldol clusters contained recurring alkoxy/benzyloxy, nitro/electron-withdrawing, sulfonyl/pyridyl and donor-substituted aldehyde subsets, together with one larger heterogeneous cluster. "
        "The cobalt failures formed a broad zero-yield collection rather than several statistically stable mechanistic families. These labels are descriptive annotations, not mechanistic assignments.",
    )

    set_paragraph(
        find_paragraph(doc, "The main experiment asks a narrow question"),
        "Changing only the acquisition representation altered the boundary regions recovered under a fixed budget. "
        "The two primary weighted-iterative-CVT contrasts showed higher absolute final-budget boundary coverage in the richer decision space than in the matched 2D space, and the bootstrap intervals for their paired differences remained positive across 20 repeats.",
    )

    image_paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph._p.xpath(".//a:blip")]
    replace_picture(image_paragraphs[0], FIG1, 6.1)
    set_caption(
        find_paragraph(doc, "Figure 1."),
        "Figure 1. ",
        "Absolute final-budget boundary coverage for the two primary paired contrasts. Bars show means over 20 repeats and error bars show bootstrap 95% confidence intervals. Within each dataset, 2D and richer-space campaigns used the same weighted iterative CVT policy, initial set and budget; evaluation was performed in the fixed full-space geometry.",
    )
    set_paragraph(
        find_paragraph(doc, "The cobalt benchmark gave the largest effect sizes"),
        "For cobalt weighted iterative CVT, mean final-budget boundary coverage increased from 0.1633 in 2D to 0.6500 in Tanimoto space. "
        "The paired increase was 0.4867 (95% CI 0.4567-0.5200). Full-space weighted iterative CVT increased coverage from the same 2D baseline to 0.4233, a paired gain of 0.2600. "
        "Because the cobalt landscape contains only 60 labelled entries and the final budget includes 24 of them, these effect sizes are interpreted as a small-data stress case rather than as a general budget-scaling result.",
    )
    set_paragraph(
        find_paragraph(doc, "Aldol provides the broader chemistry-facing support"),
        "Aldol provides the broader substrate-level test. Weighted iterative CVT increased mean boundary coverage from 0.0568 in 2D to 0.1861 in the full descriptor space, giving a paired gain of 0.1293 (95% CI 0.1236-0.1355). "
        "For conventional CVT, coverage increased from 0.0725 to 0.1908, a paired gain of 0.1183. These smaller effects are supported by a larger labelled landscape and can be connected to recurring aldehyde subsets.",
    )

    result_table = next(table for table in doc.tables if table.cell(0, 0).text.strip() == "Dataset" and "Sampler" in table.cell(0, 1).text)
    rebuild_result_table(result_table)
    set_caption(
        find_paragraph(doc, "Table 1."),
        "Table 1. ",
        "Absolute final-budget boundary coverage and paired richer-space-minus-2D differences for selected positive CVT-family contrasts. Each estimate uses 20 paired repeats. The complete method-by-space contrast matrix is reported in the Supporting Information.",
    )
    set_paragraph(
        find_paragraph(doc, "These gains should not be read as a universal ranking"),
        "The positive contrasts in Table 1 do not constitute a universal ranking of representations. The complete contrast matrix also contains neutral and negative richer-space effects for other policies. "
        "Moreover, the full-space aldol campaigns that improved boundary coverage reduced failure recall and failure-cluster coverage relative to their 2D baselines. In cobalt, Tanimoto acquisition produced the largest boundary-coverage gain but a less balanced failure-cluster profile than full-space weighted iterative CVT. "
        "Decision-space choice therefore changes which acquisition objective is emphasized.",
    )
    set_paragraph(
        find_paragraph(doc, "The most important manuscript-facing point"),
        "The additional aldol boundary points were not uniformly distributed across the labelled landscape. Recurring subsets included alkoxy- and benzyloxy-substituted aryl aldehydes, nitro or otherwise electron-withdrawing aryl aldehydes, sulfonyl or pyridyl aldehydes and donor-substituted anilino aldehydes. "
        "A larger residual cluster combined heteroaryl, halogenated and trifluoromethyl-containing aldehydes and is therefore treated as heterogeneous rather than as a single chemical family. "
        "This distinction limits the interpretation to structured substrate subsets without assigning a shared mechanism to every clustered failure.",
    )

    image_paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph._p.xpath(".//a:blip")]
    replace_picture(image_paragraphs[1], FIG2, 5.4)
    set_caption(
        find_paragraph(doc, "Figure 2."),
        "Figure 2. ",
        "Exact paired-miss examples from the aldol replay. Rows were retained only when selected by the full-space weighted-iterative-CVT replay and missed by the paired 2D replay under the same initial set and final budget. The examples provide row-level chemical checks; the repeated statistical result is reported in Figure 1 and Table 1. The full 2D location index is provided in the Supporting Information.",
    )
    remove_paragraph(find_paragraph(doc, "The full structure galleries and the larger experiment-package panel"))

    chemistry_table = next(table for table in doc.tables if table.cell(0, 0).text.strip() == "Dataset" and "Cluster" in table.cell(0, 1).text)
    chemistry_table.cell(5, 2).text = "Heterogeneous heteroaryl, halogenated and CF3 aldehydes"
    chemistry_table.cell(6, 2).text = "Heterogeneous condition-1 zero-yield alcohols"
    style_table(chemistry_table, [0.7, 0.55, 3.1, 0.5, 1.1])
    set_caption(
        find_paragraph(doc, "Table 2."),
        "Table 2. ",
        "Descriptive failure-region labels used for chemical interpretation. Labels summarize clustered entries but do not imply a shared mechanism.",
    )
    set_paragraph(
        find_paragraph(doc, "The cobalt structures are treated more cautiously"),
        "The cobalt failures are interpreted more cautiously. The largest cluster contains 32 condition-1 zero-yield alcohols spanning strained, phenoxy, pyridyl, fluorinated and polar examples. "
        "Its breadth does not support a single mechanistic-family assignment, and the remaining small or singleton clusters are used only as illustrative examples. "
        "Cobalt therefore supports the decision-space contrast numerically, while the stronger chemistry-family interpretation rests on the recurring aldol subsets.",
    )
    set_paragraph(
        find_paragraph(doc, "The main paired contrast fixes one boundary definition"),
        "The primary analysis used a 25th-percentile boundary definition and fixed failure-cluster counts. In the separate sensitivity implementation, full-space acquisition remained ahead of 2D for boundary coverage across all tested boundary quantiles and cluster counts in aldol and cobalt. "
        "The directional full-minus-2D difference ranged from 0.0433 to 0.1302 for aldol and from 0.1214 to 0.3167 for cobalt, with positive lower confidence bounds. "
        "Because this stress test used the approximation described in Experimental, these ranges support robustness of direction rather than calibration of the primary effect size.",
    )
    set_paragraph(
        find_paragraph(doc, "The same sensitivity analysis also gives a useful warning"),
        "The sensitivity analysis also exposed an endpoint tradeoff. In cobalt, Tanimoto acquisition improved failure recall and failure-cluster coverage under the stress-test implementation but produced lower boundary coverage than the 2D baseline across the tested quantiles. "
        "The representation that best samples representative failures may therefore differ from the representation that most densely covers the compatibility boundary.",
    )
    set_paragraph(
        find_paragraph(doc, "The detailed robustness ranges are reported"),
        "Full sensitivity tables, parameter grids and implementation details are reported in the Supporting Information so that the primary paired estimates remain distinguishable from the directional stress test.",
    )
    set_paragraph(
        find_paragraph(doc, "Drawable Buchwald-Hartwig is retained"),
        "External control reveals a different metric tradeoff",
    )
    set_paragraph(
        find_paragraph(doc, "A drawable Buchwald-Hartwig control was added"),
        "The drawable Buchwald-Hartwig control tested whether the same workflow could retain reaction-component structures in an external labelled dataset. In the focused FPS-style contrast, 2D gave slightly higher boundary coverage than the 32D, 64D, full and Tanimoto spaces, whereas richer spaces increased failure-cluster coverage from 0.9125 to 1.0000. "
        "This mixed result demonstrates that the benchmark can identify both advantages and disadvantages of richer decision spaces.",
    )
    set_paragraph(
        find_paragraph(doc, "The representative Buchwald-Hartwig structures"),
        "The Buchwald-Hartwig control is therefore interpreted as an external scope check rather than as evidence that one representation dominates all endpoints.",
    )
    set_paragraph(
        find_paragraph(doc, "Taken together, the results support"),
        "A 2D reaction-scope map can remain an effective interface for inspection and communication without being validated automatically as an acquisition space. "
        "Under paired budgets, changing acquisition geometry altered which near-boundary regions were observed. In aldol, part of this difference corresponded to recurring low-conversion substrate subsets rather than to isolated projection outliers.",
    )
    set_paragraph(
        find_paragraph(doc, "This framing also avoids an overclaim"),
        "The data do not establish full descriptor or Tanimoto space as universally superior. Full-space aldol acquisition improved near-boundary coverage while reducing failure recall and failure-cluster coverage; cobalt Tanimoto acquisition maximized boundary coverage in the primary weighted-iterative-CVT contrast but recovered a less balanced set of failure clusters than cobalt full-space acquisition. "
        "These disagreements show why reaction-scope acquisition should be evaluated with multiple boundary-aware endpoints.",
    )
    set_paragraph(
        find_paragraph(doc, "For practical reaction-scope studies"),
        "For practical reaction-scope studies, reports should identify the representation used for acquisition, the independent geometry used for evaluation and the endpoint being optimized. "
        "Boundary coverage, failure recall and failure-cluster diversity should be reported separately, and disagreements should be interpreted in chemical terms rather than collapsed into one aggregate score [[CITE:36,41,45]].",
    )
    set_paragraph(
        find_paragraph(doc, "Several limitations define the current scope"),
        "This study has three principal limitations. First, cobalt is a 60-entry small-data case with a final budget of 24 and cannot establish budget scaling by itself. The larger aldol landscape therefore carries the broader substrate-level argument.",
    )
    set_paragraph(
        find_paragraph(doc, "Second, the failure-region labels"),
        "Second, the failure-region labels are descriptive annotations derived from full-space clusters and simple structure tags. They demonstrate recurring substrate subsets but do not establish a mechanistic origin. The heterogeneous aldol and cobalt clusters are reported explicitly rather than forced into uniform family labels.",
    )
    set_paragraph(
        find_paragraph(doc, "Third, the robustness analysis"),
        "Third, the sensitivity analysis used an approximate weighted-iterative-CVT implementation and is restricted to directional robustness. The quantitative claim rests on the primary r20 paired contrast. The external Buchwald-Hartwig and broader reaction scans are supporting controls and do not remove the need for further prospective or experimental validation.",
    )
    set_paragraph(
        find_paragraph(doc, "A reaction-scope map can be visually useful"),
        "Visualization and acquisition are related but distinct components of reaction-scope mapping. Across paired retrospective campaigns, changing only the decision space changed the boundary regions recovered under a fixed budget. "
        "Reporting absolute coverage together with paired differences and chemistry-facing examples revealed both gains and endpoint tradeoffs. The practical conclusion is therefore not to abandon 2D maps, but to validate acquisition and boundary-discovery claims in an independently defined decision and evaluation geometry.",
    )
    set_paragraph(
        find_paragraph(doc, "The Supporting Information contains"),
        "The Supporting Information provides the complete method-by-space contrast matrix, per-repeat outputs, exact paired-miss records, full aldol and cobalt structure galleries, the 2D location index, boundary-quantile and failure-cluster sensitivity tables, and the drawable Buchwald-Hartwig control.",
    )
    set_paragraph(
        find_paragraph(doc, "Author contributions will be finalized"),
        "Author contributions must be confirmed by both authors and entered using the CRediT taxonomy before submission. Required roles include conceptualization, data curation, software, formal analysis, visualization, writing-original draft, writing-review and editing, supervision and project administration; roles must be assigned only after author confirmation.",
    )
    set_paragraph(
        find_paragraph(doc, "The analysis scripts, configuration files"),
        "The analysis scripts, configuration files, per-repeat outputs, derived tables and figure inputs are maintained in the project archive. A public repository URL, version tag and archival DOI must be inserted here after deposition and before submission. "
        "The original aldol and cobalt datasets derive from the cited ScopeMap study; the external Buchwald-Hartwig source and all preprocessing steps must be identified in the accompanying repository and Supporting Information.",
    )
    set_paragraph(
        find_paragraph(doc, "Funding, institutional support"),
        "Funding, institutional support, contributor acknowledgements and the journal-required AI-tool disclosure must be completed and approved by the authors before submission. No funding source or contributor has been inferred in this revision.",
    )

    # Complete author lists required by RSC reference style.
    set_paragraph(
        find_paragraph(doc, "S. Szymkuc, E. P. Gajewska"),
        "S. Szymkuc, E. P. Gajewska, T. Klucznik, K. Molga, P. Dittwald, M. Startek, M. Bajczyk and B. A. Grzybowski, Angew. Chem. Int. Ed., 2016, 55, 5904-5937, DOI: 10.1002/anie.201506101.",
    )
    set_paragraph(
        find_paragraph(doc, "M. Christensen, L. P. E. Yunker"),
        "M. Christensen, L. P. E. Yunker, P. Shiri, T. Zepel, P. L. Prieto, S. Grunert, F. Bork and J. E. Hein, Chem. Sci., 2021, 12, 15473-15490, DOI: 10.1039/D1SC04588A.",
    )

    # Keep all captions black, non-italic and readable.
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Caption":
            for run in paragraph.runs:
                run.italic = False
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = Pt(9)

    # Prevent headings and captions from being stranded at page bottoms.
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading") or paragraph.style.name == "Caption":
            paragraph.paragraph_format.keep_with_next = True

    doc.core_properties.title = "Visualization Space Is Not Validation Space"
    doc.core_properties.subject = "Revised Chemical Science manuscript"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    revise()
